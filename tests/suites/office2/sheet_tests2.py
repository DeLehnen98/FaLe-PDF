"""Spreadsheet part of the Office 2 tests (imported by test_office2.py)."""
import os, json, base64, time, zipfile, re

A = "document.querySelector('.sheet-app')._sheetApi"
FILES_JS = """async () => { const db = await new Promise((r, j) => { const q = indexedDB.open('falepdf'); q.onsuccess = () => r(q.result); q.onerror = () => j(q.error); });
  const res = await new Promise(r => { const t = db.transaction('docs'); const g = t.objectStore('docs').getAll(); g.onsuccess = () => r(g.result.map(d => ({id: d.id, name: d.name, kind: d.kind || 'pdf'}))); });
  db.close(); return res; }"""
BLOB_JS = """async (id) => { const db = await new Promise((r, j) => { const q = indexedDB.open('falepdf'); q.onsuccess = () => r(q.result); q.onerror = () => j(q.error); });
  const rec = await new Promise(r => { const t = db.transaction('files'); const g = t.objectStore('files').get(id); g.onsuccess = () => r(g.result); });
  db.close(); if (!rec) return null;
  return await new Promise(r => { const fr = new FileReader(); fr.onload = () => r(fr.result.split(',')[1]); fr.readAsDataURL(rec.blob); }); }"""


def run(p, check, step, attach_logging, URL, OUT, WORK):
    import openpyxl
    ctx = p.chromium.launch().new_context(viewport={'width': 1366, 'height': 860}, locale='de-DE', accept_downloads=True, bypass_csp=True)
    page = ctx.new_page()
    attach_logging(page, 'sheet')
    page.goto(URL)
    page.wait_for_selector('#view-start .hero')

    def go(h, wait=500):
        page.evaluate(f"location.hash = {json.dumps(h)}")
        page.wait_for_timeout(wait)

    def new_sheet():
        go('#/start', 200)
        dlg = page.locator('dialog', has_text='Änderungen speichern?')
        if dlg.count():
            dlg.locator('button[value="discard"]').click()
            page.wait_for_timeout(300)
        go('#/neu/tabelle', 300)
        page.wait_for_selector('.sheet-app', timeout=10000)
        page.wait_for_timeout(300)
        info = page.evaluate("() => [document.querySelectorAll('.sheet-app').length, location.hash, !!(document.querySelector('.sheet-app') && document.querySelector('.sheet-app')._sheetApi.sheet())]")
        if not info[2]:
            print('new_sheet state', info)

    def inp(a, v):
        page.evaluate(f"([a, v]) => {A}.input(a, v)", [a, v])

    def txt(a):
        return page.evaluate(f"a => {A}.text(a)", a)

    def val(a):
        return page.evaluate(f"a => {A}.value(a)", a)

    def sel(a):
        page.evaluate(f"a => {A}.select(a)", a)
        page.wait_for_timeout(30)

    def leave():
        page.evaluate("location.hash = '#/start'")
        page.wait_for_timeout(300)
        dlg = page.locator('dialog', has_text='Änderungen speichern?')
        if dlg.count():
            dlg.locator('button[value="discard"]').click()
            page.wait_for_timeout(300)

    def save(name):
        page.locator('[data-office-save]').click()
        page.wait_for_function("document.querySelector('#tool-screen .screen-head .sub') && /Gespeichert/.test(document.querySelector('#tool-screen .screen-head .sub').textContent)", timeout=20000)
        page.wait_for_timeout(300)
        hid = page.evaluate("location.hash").split('/')[-1].split('?')[0]
        data = base64.b64decode(page.evaluate(BLOB_JS, hid))
        path = os.path.join(WORK, name)
        open(path, 'wb').write(data)
        return path

    def shot(name):
        page.wait_for_timeout(300)
        page.screenshot(path=os.path.join(OUT, f'{name}.png'))

    def formulas():
        new_sheet()
        data = [['Name', 'Abteilung', 'Umsatz', 'Datum'], ['Anna', 'Nord', '1200', '03.01.2025'], ['Ben', 'Süd', '800', '15.02.2025'], ['Cem', 'Nord', '1500', '20.02.2025'], ['Dora', 'West', '950', '01.03.2025'], ['Emil', 'Süd', '400', '12.03.2025']]
        for r, row in enumerate(data):
            for c, v in enumerate(row):
                inp('ABCD'[c] + str(r + 1), v)
        F = {
            'F1': ('=XVERWEIS("Cem";A2:A6;C2:C6)', '1500'),
            'F2': ('=SVERWEIS("Ben";A2:C6;3;FALSCH)', '800'),
            'F3': ('=SUMMEWENNS(C2:C6;B2:B6;"Nord")', '2700'),
            'F4': ('=ZÄHLENWENN(B2:B6;"Süd")', '2'),
            'F5': ('=WENNS(C2>1000;"hoch";C2>500;"mittel";WAHR;"niedrig")', 'hoch'),
            'F6': ('=TEXTVERKETTEN(", ";WAHR;A2:A3)', 'Anna, Ben'),
            'F7': ('=SVERWEIS("Zoe";A2:C6;3;FALSCH)', '#NV'),
            'F8': ('=WENNFEHLER(1/0;"Fehler")', 'Fehler'),
            'F9': ('="a"+1', '#WERT!'),
            'F10': ('=MONAT(D3)', '2'),
            'F11': ('=NETTOARBEITSTAGE(D2;D3)', '31'),
            'F12': ('=RUNDEN(MITTELWERT(C2:C6);1)', '970'),
            'F13': ('=SUMMENPRODUKT((B2:B6="Süd")*C2:C6)', '1200'),
            'F14': ('=INDEX(A2:A6;VERGLEICH(MAX(C2:C6);C2:C6;0))', 'Cem'),
            'F15': ('=TEXT(1234,5;"#.##0,00 €")', '1.234,50 €'),
            'F16': ('=GROSS2("hallo welt")', 'Hallo Welt'),
            'F17': ('=RUNDEN(ZW(5%/12;12;-100);2)', '1227,89'),
            'F18': ('=LET(x;3;x*x)', '9'),
        }
        for a, (f, _) in F.items():
            inp(a, f)
        page.wait_for_timeout(200)
        got = {a: txt(a) for a in F}
        bad = {a: (got[a], e) for a, (f, e) in F.items() if got[a] != e}
        check('German formulas incl. XVERWEIS/WENNS/TEXTVERKETTEN/LET + error values', not bad, str(bad))
        # dynamic arrays with spill
        inp('H1', '=SORTIEREN(EINDEUTIG(B2:B6))')
        inp('I1', '=FILTER(A2:A6;C2:C6>900)')
        inp('J1', '=SEQUENZ(3;2;10;5)')
        inp('L1', '=ANZAHL2(I1#)')
        page.wait_for_timeout(200)
        check('EINDEUTIG+SORTIEREN spill', [txt(f'H{i}') for i in range(1, 4)] == ['Nord', 'Süd', 'West'])
        check('FILTER spill', [txt(f'I{i}') for i in range(1, 5)] == ['Anna', 'Cem', 'Dora', ''])
        check('SEQUENZ 3x2', [txt('J1'), txt('K1'), txt('J3'), txt('K3')] == ['10', '15', '30', '35'])
        check('spill reference I1#', txt('L1') == '3')
        inp('I3', 'blockiert')
        page.wait_for_timeout(100)
        check('#ÜBERLAUF! when spill area is blocked', txt('I1') == '#ÜBERLAUF!', txt('I1'))
        inp('I3', '')
        # circular reference + cross-sheet + names
        inp('M1', '=M2+1'); inp('M2', '=M1+1')
        page.wait_for_timeout(100)
        check('circular reference detected', txt('M1') in ('#ZIRKEL!', '0') and txt('M2') in ('#ZIRKEL!', '0'), txt('M1'))
        inp('M1', ''); inp('M2', '')
        page.evaluate(f"() => {{ const wb = {A}.wb(); wb.names = (wb.names || []).concat([{{ name: 'Umsätze', ref: 'Tabelle1!$C$2:$C$6' }}]); {A}.recalcAll(); }}")
        inp('N1', '=SUMME(Umsätze)')
        page.wait_for_timeout(100)
        check('named range in formula', txt('N1') == '4850', txt('N1'))
        # dependent recalculation
        inp('C2', '2200')
        page.wait_for_timeout(100)
        check('dependency recalculation (SUMMEWENNS + name + spill)', txt('F3') == '3700' and txt('N1') == '5850' and txt('I1') == 'Anna')
        # absolute refs + fill + Ctrl+D
        inp('P1', '2'); inp('Q1', '=$P$1*C2')
        sel('Q1:Q5')
        page.evaluate(f"{A}.action('fill-down')") if False else None
        page.locator('.sh-scroll').focus()
        page.keyboard.press('Control+d')
        page.wait_for_timeout(150)
        check('Ctrl+D fills formulas with relative/absolute refs', txt('Q3') == '3000' and txt('Q5') == '800', f'{txt("Q3")} {txt("Q5")}')
        # autofill series
        inp('S1', 'Mo'); inp('S2', 'Di')
        sel('S1:S2')
        page.evaluate(f"{A}.fill('down', 5)")
        inp('T1', 'Januar')
        sel('T1')
        page.evaluate(f"{A}.fill('down', 3)")
        inp('U1', '01.01.2025'); inp('U2', '01.02.2025')
        sel('U1:U2')
        page.evaluate(f"{A}.fill('down', 2)")
        inp('V1', '5'); inp('V2', '10')
        sel('V1:V2')
        page.evaluate(f"{A}.fill('down', 3)")
        inp('W1', 'Q1')
        sel('W1')
        page.evaluate(f"{A}.fill('right', 3)")
        page.wait_for_timeout(150)
        check('autofill weekdays', [txt(f'S{i}') for i in range(3, 8)] == ['Mi', 'Do', 'Fr', 'Sa', 'So'], str([txt(f'S{i}') for i in range(3, 8)]))
        check('autofill months', [txt(f'T{i}') for i in range(2, 5)] == ['Februar', 'März', 'April'])
        check('autofill monthly dates', [txt('U3'), txt('U4')] == ['01.03.2025', '01.04.2025'], f'{txt("U3")} {txt("U4")}')
        check('autofill linear numbers', [txt(f'V{i}') for i in range(3, 6)] == ['15', '20', '25'])
        check('autofill Q1..Q4 to the right', [txt('X1'), txt('Y1'), txt('Z1')] == ['Q2', 'Q3', 'Q4'])
        # undo / redo
        page.evaluate(f"{A}.undo()")
        page.wait_for_timeout(100)
        u = txt('X1')
        page.evaluate(f"{A}.redo()")
        page.wait_for_timeout(100)
        check('undo / redo of autofill', u == '' and txt('X1') == 'Q2')
        # clipboard: external TSV paste, internal copy + transpose paste
        sel('AA1')
        page.evaluate(f"{A}.paste('Stadt\\tEinwohner\\nKöln\\t1.084.000\\nBonn\\t336.000')")
        page.wait_for_timeout(100)
        check('paste external TSV (German number parsing)', txt('AA2') == 'Köln' and val('AB2') == 1084000, f'{txt("AA2")} {val("AB2")}')
        sel('AA1:AB3')
        page.evaluate(f"{A}.copy()")
        sel('AD1')
        page.evaluate(f"{A}.pasteInternal({{ transpose: true }})")
        page.wait_for_timeout(100)
        check('paste special: transpose', [txt('AD1'), txt('AE1'), txt('AF1'), txt('AD2')] == ['Stadt', 'Köln', 'Bonn', 'Einwohner'])
        sel('F1:F3')
        page.evaluate(f"{A}.copy()")
        sel('AH1')
        page.evaluate(f"{A}.pasteInternal({{ what: 'values' }})")
        page.wait_for_timeout(100)
        check('paste special: values only', page.evaluate(f"() => {{ const c = {A}.sheet().cells.get(33); return !!c && !c.f && c.v === 1500; }}"))
        shot('desktop_sheet_formulas_light')
        leave()
    step('sheet formulas, spill, fill, clipboard', formulas)

    def data_tools():
        new_sheet()
        rows = [['Region', 'Produkt', 'Menge', 'Preis'], ['Nord', 'Äpfel', 10, 2.5], ['Süd', 'Birnen', 5, 3], ['Nord', 'Birnen', 7, 3], ['West', 'Äpfel', 3, 2.5], ['Süd', 'Äpfel', 8, 2.5], ['Nord', 'Äpfel', 10, 2.5]]
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                inp('ABCD'[c] + str(r + 1), str(v).replace('.', ','))
        # pivot
        page.evaluate(f"{A}.pivot({{ src: 'Tabelle1', range: 'A1:D7', def: {{ rows: [0], cols: [1], vals: [{{ c: 2, agg: 'sum' }}] }} }})")
        page.wait_for_timeout(400)
        names = page.evaluate(f"{A}.wb().sheets.map(s => s.name)")
        check('pivot table on new sheet', any('Pivot' in n for n in names), str(names))
        grid = page.evaluate(f"() => {{ const sh = {A}.sheet(); const out = []; for (let r = 0; r <= sh.maxR; r++) {{ const row = []; for (let c = 0; c <= sh.maxC; c++) row.push({A}.text(String.fromCharCode(65 + c) + (r + 1))); out.push(row); }} return out; }}")
        flat = [x for row in grid for x in row]
        nord = [row for row in grid if row and row[0] == 'Nord']
        check('pivot sums (Nord: Äpfel 20, Birnen 7)', nord and '20' in nord[0] and '7' in nord[0], str(grid))
        check('pivot grand total 43', '43' in flat, str(grid[-1] if grid else ''))
        page.evaluate(f"{A}.switchSheet(0)")
        page.wait_for_timeout(200)
        # multi-column sort: Region asc, Menge desc
        page.evaluate(f"{A}.sortRange('A2:D7', [{{ c: 0, asc: true }}, {{ c: 2, asc: false }}])")
        page.wait_for_timeout(150)
        check('multi-column sort', [txt(f'A{i}') + txt(f'C{i}') for i in range(2, 8)] == ['Nord10', 'Nord10', 'Nord7', 'Süd8', 'Süd5', 'West3'], str([txt(f'A{i}') + txt(f'C{i}') for i in range(2, 8)]))
        # custom number formats
        inp('F1', '1234,5'); inp('F2', '0,256'); inp('F3', '01.03.2025'); inp('F4', '-42')
        page.evaluate("""() => { const api = document.querySelector('.sheet-app')._sheetApi; const sh = api.sheet(); const set = (k, f) => { const c = sh.cells.get(k); sh.cells.set(k, { ...c, s: { ...(c.s || {}), fmt: f } }); }; set(5, '#,##0.00 "€"'); set(16384 + 5, '0.0%'); set(2 * 16384 + 5, 'dddd, dd. mmmm yyyy'); set(3 * 16384 + 5, '#,##0;[Red]-#,##0'); api.recalcAll(); }""")
        page.wait_for_timeout(150)
        check('custom number formats (currency, percent, long date, negative red)', [txt('F1'), txt('F2'), txt('F3'), txt('F4')] == ['1.234,50 €', '25,6 %', 'Samstag, 01. März 2025', '-42'], str([txt('F1'), txt('F2'), txt('F3'), txt('F4')]))
        # dedupe via dialog
        sel('A1:D7')
        page.locator('.rb-tab[data-rb="data"]').click()
        page.locator('[data-a="dedupe"]').click()
        page.wait_for_timeout(200)
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(300)
        check('remove duplicates', page.evaluate(f"{A}.sheet().maxR") == 5 or txt('A7') == '', f'{txt("A6")}|{txt("A7")}')
        # protection: locked cell refuses input
        page.evaluate(f"{A}.action('protect')")
        page.wait_for_timeout(200)
        page.locator('#sh-addr').fill('B2'); page.locator('#sh-addr').press('Enter'); page.wait_for_timeout(100)
        page.keyboard.type('Geändert'); page.keyboard.press('Enter')
        page.wait_for_timeout(200)
        check('sheet protection blocks edits', txt('B2') != 'Geändert')
        page.evaluate(f"{A}.action('protect')")
        # freeze at a cell + hidden rows
        page.evaluate(f"{A}.freeze(1, 1)")
        page.wait_for_timeout(200)
        check('freeze panes at B2', page.locator('.sh-frozen').is_visible() and page.locator('.sh-body .sh-fc').count() >= 1)
        page.evaluate(f"() => {{ const sh = {A}.sheet(); sh.hiddenRows.add(4); sh.rowH.set(1, 40); sh.colW[1] = 160; sh.tabColor = '#15803D'; sh.merges = (sh.merges || []).concat([{{ r1: 9, c1: 0, r2: 9, c2: 2 }}]); {A}.recalcAll(); }}")
        inp('A10', 'Verbunden')
        page.evaluate(f"() => {{ const sh = {A}.sheet(); const c = sh.cells.get(9 * 16384); sh.cells.set(9 * 16384, {{ ...c, s: {{ b: true, i: true, font: 'Georgia', size: 14, color: '#B91C1C', fill: '#FFF4CC', align: 'center', valign: 'center', wrap: true, rot: 90, bd: {{ t: 'thin', b: 'double' }} }} }}); {A}.recalcAll(); }}")
        page.wait_for_timeout(150)
        # data validation + autofilter
        page.evaluate(f"() => {{ const sh = {A}.sheet(); sh.dv = (sh.dv || []).concat([{{ range: 'E2:E7', type: 'list', list: ['ja', 'nein'] }}]); sh.autoFilter = {{ range: 'A1:D6', filters: {{}} }}; {A}.recalcAll(); }}")
        inp('G1', '=XVERWEIS("West";A2:A6;C2:C6)')
        inp('G2', '=SORTIEREN(C2:C6;1;-1)')
        shot('desktop_sheet_tools_light')
        path = save('daten.xlsx')
        wb = openpyxl.load_workbook(path)
        ws = wb.worksheets[0]
        check('openpyxl: workbook opens, sheets incl. pivot sheet', len(wb.sheetnames) >= 2, str(wb.sheetnames))
        check('openpyxl: freeze panes B2', ws.freeze_panes == 'B2', str(ws.freeze_panes))
        check('openpyxl: hidden row 5', ws.row_dimensions[5].hidden)
        check('openpyxl: row height + column width', abs((ws.row_dimensions[2].height or 0) - 21.43) < 0.5 and (ws.column_dimensions['B'].width or 0) > 18, f'{ws.row_dimensions[2].height} {ws.column_dimensions["B"].width}')
        check('openpyxl: tab colour', ws.sheet_properties.tabColor is not None and '15803D' in str(ws.sheet_properties.tabColor.rgb))
        check('openpyxl: merged cells', 'A10:C10' in [str(r) for r in ws.merged_cells.ranges])
        c = ws['A10']
        check('openpyxl: font name/size/bold/italic/colour', c.font.name == 'Georgia' and c.font.sz == 14 and c.font.b and c.font.i and 'B91C1C' in str(c.font.color.rgb), f'{c.font.name} {c.font.sz} {c.font.color.rgb if c.font.color else None}')
        check('openpyxl: fill + alignment + rotation + wrap', 'FFF4CC' in str(c.fill.fgColor.rgb) and c.alignment.horizontal == 'center' and c.alignment.vertical == 'center' and c.alignment.textRotation == 90 and c.alignment.wrap_text, f'{c.fill.fgColor.rgb} {c.alignment.horizontal} {c.alignment.textRotation}')
        check('openpyxl: borders', c.border.top.style == 'thin' and c.border.bottom.style == 'double')
        check('openpyxl: number formats', ws['F1'].number_format.replace('\\', '') in ('#,##0.00 "€"', '#,##0.00 €') and ws['F2'].number_format == '0.0%', f'{ws["F1"].number_format} {ws["F2"].number_format}')
        check('openpyxl: XLOOKUP stored with _xlfn prefix', '_xlfn.XLOOKUP' in str(ws['G1'].value), str(ws['G1'].value))
        xml = zipfile.ZipFile(path).read('xl/worksheets/sheet1.xml').decode()
        check('xlsx: dynamic array formula (t="array" + cm)', 't="array"' in xml and 'cm="1"' in xml and '_xlfn._xlws.SORT' in xml or '_xlfn.SORT' in xml)
        check('openpyxl: data validation list', any(dv.type == 'list' for dv in ws.data_validations.dataValidation))
        check('openpyxl: autofilter', ws.auto_filter.ref == 'A1:D6', str(ws.auto_filter.ref))
        check('openpyxl: defined pivot values present', any(cell.value == 43 for row in wb.worksheets[1].iter_rows() for cell in row))
        # re-open in the app
        page.reload()
        page.wait_for_selector('.sheet-app', timeout=15000)
        page.wait_for_timeout(500)
        check('re-open: values, spill, merge, freeze', txt('G1') == '3' and txt('G2') == '10' and txt('G3') == '8' and page.evaluate(f"{A}.sheet().freezeC") == 1 and len(page.evaluate(f"{A}.sheet().merges")) >= 1, f'{txt("G1")} {txt("G2")} {txt("G3")}')
        leave()
    step('sheet pivot, sort, formats, protection, xlsx round trip', data_tools)

    def sheets_tabs():
        new_sheet()
        inp('A1', '5')
        page.locator('.sh-tab').first.click(button='right')
        page.wait_for_timeout(150)
        page.locator('.menu-item', has_text='Duplizieren').click()
        page.wait_for_timeout(200)
        check('duplicate sheet', page.locator('.sh-tab').count() >= 2)
        inp('B1', "='Tabelle1'!A1*2")
        page.wait_for_timeout(100)
        check('cross-sheet reference', txt('B1') == '10', txt('B1'))
        page.locator('.sh-tab').first.click(button='right')
        page.locator('.menu-item', has_text='Registerfarbe').click()
        page.wait_for_timeout(100)
        page.locator('.menu-item', has_text='Petrol').click()
        page.wait_for_timeout(100)
        check('tab colour', page.locator('.sh-tab.has-color').count() >= 1)
        leave()
    step('sheet tabs', sheets_tabs)

    def perf():
        new_sheet()
        t = page.evaluate(f"""() => {{ const api = {A}; const sh = api.sheet(); const t0 = performance.now();
            for (let r = 0; r < 10000; r++) {{ for (let c = 0; c < 9; c++) sh.cells.set(r * 16384 + c, {{ v: r * 10 + c }}); sh.cells.set(r * 16384 + 9, {{ f: 'SUM(A' + (r + 1) + ':I' + (r + 1) + ')' }}); }}
            sh.maxR = 9999; sh.maxC = 9; const t1 = performance.now(); api.recalcAll(); const t2 = performance.now(); return [t1 - t0, t2 - t1]; }}""")
        page.wait_for_timeout(500)
        check('100 000 cells: full recalculation < 3 s', t[1] < 3000, f'{t[1]:.0f} ms')
        check('100k: row sums correct', txt('J1') == '36' and txt('J10000') == str(sum(9999 * 10 + c for c in range(9))), txt('J10000'))
        t0 = time.time()
        inp('A5000', '1000000')
        page.wait_for_timeout(50)
        dt = (time.time() - t0) * 1000
        check('100k: single edit recalculates incrementally (< 400 ms incl. round trip)', dt < 400 and txt('J5000') == str(1000000 + sum(4999 * 10 + c for c in range(1, 9))), f'{dt:.0f} ms {txt("J5000")}')
        t0 = time.time()
        for y in range(0, 200000, 20000):
            page.evaluate(f"document.querySelector('.sh-scroll').scrollTop = {y}")
            page.wait_for_timeout(16)
        page.evaluate("new Promise(r => requestAnimationFrame(() => requestAnimationFrame(r)))")
        dt = (time.time() - t0) * 1000
        check('100k: scrolling through 10 000 rows stays responsive', dt < 2500 and page.locator('.sh-body [data-cell]').count() < 3000, f'{dt:.0f} ms, {page.locator(".sh-body [data-cell]").count()} cells in DOM')
        leave()
    step('sheet performance 100k cells', perf)

    def shots():
        new_sheet()
        rows = [['Monat', 'Umsatz', 'Kosten', 'Gewinn'], ['Januar', '1200', '800', '=B2-C2'], ['Februar', '1500', '900', '=B3-C3'], ['März', '1700', '950', '=B4-C4']]
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                inp('ABCD'[c] + str(r + 1), v)
        inp('F1', '=SORTIEREN(B2:B4;1;-1)')
        page.evaluate(f"{A}.freeze(1, 1)")
        sel('B2:D4')
        shot('desktop_sheet_light')
        page.evaluate("() => { localStorage.setItem('fale-theme', 'dark'); document.documentElement.dataset.theme = 'dark'; }")
        shot('desktop_sheet_dark')
        page.evaluate("() => { localStorage.setItem('fale-theme', 'light'); document.documentElement.dataset.theme = 'light'; }")
        leave()
    step('sheet screenshots', shots)
    ctx.close()

    ctx = p.chromium.launch().new_context(viewport={'width': 390, 'height': 844}, device_scale_factor=2, is_mobile=True, has_touch=True, locale='de-DE', bypass_csp=True)
    page = ctx.new_page()
    attach_logging(page, 'sheet-phone')
    page.goto(URL)
    page.wait_for_selector('#view-start .hero')

    def phone():
        go('#/neu/tabelle', 300)
        page.wait_for_selector('.sheet-app', timeout=10000)
        page.wait_for_timeout(300)
        for a, v in [('A1', 'Monat'), ('B1', 'Umsatz'), ('A2', 'Januar'), ('B2', '1200'), ('A3', 'Februar'), ('B3', '1500'), ('C1', '=SUMME(B2:B3)')]:
            page.evaluate(f"([a, v]) => {A}.input(a, v)", [a, v])
        sw = page.evaluate('document.documentElement.scrollWidth')
        check('phone sheet: no horizontal overflow', sw <= 390, str(sw))
        page.wait_for_timeout(300)
        page.screenshot(path=os.path.join(OUT, 'phone_sheet_light.png'))
        page.evaluate("() => { localStorage.setItem('fale-theme', 'dark'); document.documentElement.dataset.theme = 'dark'; }")
        page.wait_for_timeout(300)
        page.screenshot(path=os.path.join(OUT, 'phone_sheet_dark.png'))
        page.evaluate("() => { localStorage.setItem('fale-theme', 'light'); document.documentElement.dataset.theme = 'light'; }")
    step('sheet phone', phone)
    ctx.close()
