"""Office 2 tests: document editor (page layout, styles, formatting, lists, tables, images, track changes, mail merge,
find options, versions, shortcuts) and spreadsheet (functions, spill, fill, clipboard, pivot, formats, sort/filter,
protection, xlsx round trip via openpyxl, 100k-cell performance). Server: 127.0.0.1:9202 (private copy, see sync.sh)."""
import os, sys, io, json, base64, subprocess, traceback, shutil, time, zipfile, re, datetime
from playwright.sync_api import sync_playwright

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'out')
WORK = os.path.join(HERE, 'work')
shutil.rmtree(WORK, ignore_errors=True)
os.makedirs(OUT, exist_ok=True)
os.makedirs(WORK, exist_ok=True)
URL = os.environ.get('OFFICE_URL', 'http://127.0.0.1:9202/')
ONLY = sys.argv[1] if len(sys.argv) > 1 else 'all'
results, console_problems = [], []


def check(name, cond, info=''):
    results.append((name, bool(cond), info))
    print(('PASS ' if cond else 'FAIL ') + name + (f'  [{info}]' if info and not cond else ''), flush=True)


def step(name, fn):
    try:
        fn()
    except Exception as e:
        results.append((name + ' (exception)', False, str(e)[:300]))
        print('FAIL ' + name + ' EXC ' + str(e)[:800], flush=True)
        traceback.print_exc()


def attach_logging(page, label):
    page.on('console', lambda m: console_problems.append(f'{label} [{m.type}] {m.text}') if m.type in ('error', 'warning') and 'favicon' not in m.text else None)
    page.on('pageerror', lambda e: console_problems.append(f'{label} [pageerror] {e}'))


FILES_JS = """async () => { const db = await new Promise((r, j) => { const q = indexedDB.open('falepdf'); q.onsuccess = () => r(q.result); q.onerror = () => j(q.error); });
  const res = await new Promise(r => { const t = db.transaction('docs'); const g = t.objectStore('docs').getAll(); g.onsuccess = () => r(g.result.map(d => ({id: d.id, name: d.name, kind: d.kind || 'pdf'}))); });
  db.close(); return res; }"""
BLOB_JS = """async (id) => { const db = await new Promise((r, j) => { const q = indexedDB.open('falepdf'); q.onsuccess = () => r(q.result); q.onerror = () => j(q.error); });
  const rec = await new Promise(r => { const t = db.transaction('files'); const g = t.objectStore('files').get(id); g.onsuccess = () => r(g.result); });
  db.close(); if (!rec) return null;
  return await new Promise(r => { const fr = new FileReader(); fr.onload = () => r(fr.result.split(',')[1]); fr.readAsDataURL(rec.blob); }); }"""


def files(page):
    return page.evaluate(FILES_JS)


def save_blob(page, doc_id, name):
    data = base64.b64decode(page.evaluate(BLOB_JS, doc_id))
    path = os.path.join(WORK, name)
    open(path, 'wb').write(data)
    return path


def goto_hash(page, h, wait=500):
    page.evaluate(f"location.hash = {json.dumps(h)}")
    page.wait_for_timeout(wait)


def shot(page, name):
    page.wait_for_timeout(300)
    page.screenshot(path=os.path.join(OUT, f'{name}.png'))


def set_theme(page, t):
    page.evaluate("t => { localStorage.setItem('fale-theme', t); document.documentElement.dataset.theme = t; }", t)
    page.wait_for_timeout(150)


def leave(page):
    page.evaluate("location.hash = '#/start'")
    page.wait_for_timeout(300)
    dlg = page.locator('dialog', has_text='Änderungen speichern?')
    if dlg.count():
        dlg.locator('button[value="discard"]').click()
        page.wait_for_timeout(400)


def wait_saved(page, timeout=20000):
    page.wait_for_function("document.querySelector('#tool-screen .screen-head .sub') && /Gespeichert/.test(document.querySelector('#tool-screen .screen-head .sub').textContent)", timeout=timeout)


def save_doc(page):
    page.locator('[data-office-save]').click()
    wait_saved(page)
    page.wait_for_timeout(300)
    hid = page.evaluate("location.hash").split('/')[-1].split('?')[0]
    return hid


def menu(page, text):
    page.locator('.menu-item', has_text=text).first.click()
    page.wait_for_timeout(200)


# ---------- writer helpers ----------
W = "document.querySelector('.writer')._writerApi"
SEL_JS = """([text, start, len]) => { const ed = document.querySelector('.doc-page'); const tw = document.createTreeWalker(ed, NodeFilter.SHOW_TEXT); let n;
  while ((n = tw.nextNode())) { const i = n.nodeValue.indexOf(text); if (i >= 0) { const r = document.createRange(); r.setStart(n, i + (start || 0)); r.setEnd(n, i + (start || 0) + (len == null ? text.length : len)); const s = getSelection(); s.removeAllRanges(); s.addRange(r); ed.focus(); return true; } } return false; }"""
CARET_JS = """([text, end]) => { const ed = document.querySelector('.doc-page'); const tw = document.createTreeWalker(ed, NodeFilter.SHOW_TEXT); let n;
  while ((n = tw.nextNode())) { const i = n.nodeValue.indexOf(text); if (i >= 0) { const r = document.createRange(); r.setStart(n, end ? i + text.length : i); r.collapse(true); const s = getSelection(); s.removeAllRanges(); s.addRange(r); ed.focus(); return true; } } return false; }"""


def wsel(page, text, start=0, length=None):
    assert page.evaluate(SEL_JS, [text, start, length]), f'text not found: {text}'
    page.wait_for_timeout(60)


def wcaret(page, text, end=True):
    assert page.evaluate(CARET_JS, [text, end]), f'text not found: {text}'
    page.wait_for_timeout(60)


def set_html(page, html):
    page.evaluate("h => { const ed = document.querySelector('.doc-page'); ed.innerHTML = h; ed.dispatchEvent(new Event('input', { bubbles: true })); }", html)
    page.wait_for_timeout(400)


def rb(page, tab):
    page.locator(f'.rb-tab[data-rb="{tab}"]').click()
    page.wait_for_timeout(100)


def new_doc(page):
    goto_hash(page, '#/neu/dokument', 300)
    page.wait_for_selector('.doc-page', timeout=10000)
    page.wait_for_timeout(400)


def docx_xml(path, part='word/document.xml'):
    with zipfile.ZipFile(path) as z:
        return z.read(part).decode('utf-8')


# =====================================================================
def writer_tests(p):
    from docx import Document
    from docx.shared import Pt
    ctx = p.chromium.launch().new_context(viewport={'width': 1366, 'height': 860}, locale='de-DE', accept_downloads=True, bypass_csp=True)
    page = ctx.new_page()
    attach_logging(page, 'writer')
    page.goto(URL)
    page.wait_for_selector('#view-start .hero')

    def layout():
        new_doc(page)
        check('desktop opens in page layout view', page.locator('.doc-sheet.pv').count() == 1 and page.locator('.pv-page').count() >= 1)
        paras = ''.join(f'<p>Absatz {i}: ' + ('Lorem ipsum dolor sit amet, consectetur adipiscing elit. ' * 7) + '</p>' for i in range(16))
        set_html(page, '<h1>Bericht</h1>' + paras + '<hr data-page-break><h2>Anhang</h2><p>Letzte Seite.</p>')
        page.wait_for_timeout(600)
        n = page.locator('.pv-page').count()
        check('long document paginates into several pages', n >= 3, str(n))
        # no text line may be placed inside a page gap / margin
        bad = page.evaluate("""() => { const ed = document.querySelector('.doc-page'); const base = ed.getBoundingClientRect().top; const ph = parseFloat(getComputedStyle(document.querySelector('.doc-sheet')).getPropertyValue('--page-h')); const stride = ph + 22;
          const mt = parseFloat(getComputedStyle(ed).getPropertyValue('--pad-t')), mb = parseFloat(getComputedStyle(ed).getPropertyValue('--pad-b')); let bad = 0, lines = 0;
          ed.querySelectorAll('p,h1,h2').forEach(p => { const r = document.createRange(); r.selectNodeContents(p); for (const q of r.getClientRects()) { if (q.height < 2) continue; lines++; const y0 = q.top - base, y1 = q.bottom - base; const k = Math.floor(y0 / stride); const top = k * stride + mt - 1, bot = k * stride + ph - mb + 1; if (y0 < top || y1 > bot) bad++; } }); return [bad, lines]; }""")
        check('every line lies inside the text area of a page', bad[0] == 0 and bad[1] > 50, str(bad))
        pg_anhang = page.evaluate(f"{W}.pv.pageOf(document.querySelector('.doc-page h2'))")
        check('hard page break starts a new page', pg_anhang == n, f'{pg_anhang} / {n}')
        check('status bar shows page x of y', f'von {n}' in page.locator('[data-pageinfo]').inner_text())
        # header/footer with page numbers on every page
        page.evaluate(f"{W}.setHf({{ header: 'Kopfzeile Test', headerAlign: 'left', pageNumbers: 'footer-center', pageFormat: 'n-of-m' }})")
        page.wait_for_timeout(500)
        hf_texts = page.locator('.pv-hf.bottom').all_inner_texts()
        check('page numbers on following pages', f'Seite 2 von {n}' in hf_texts and f'Seite {n} von {n}' in hf_texts, str(hf_texts[:3]))
        check('page 1 header via header button', 'Kopfzeile Test' in page.locator('.doc-hf.top').inner_text() and page.locator('.pv-hf.top').count() == n - 1)
        # zoom
        rb(page, 'view')
        page.locator('[data-x="zoom-fit"]').click()
        page.wait_for_timeout(300)
        w = page.locator('.doc-sheet').bounding_box()['width']
        sw = page.locator('.writer-scroll').bounding_box()['width']
        check('zoom: fit page width', abs(w - (sw - 40)) < 30, f'{w} vs {sw}')
        page.locator('[data-zoom]').select_option('0.5')
        page.wait_for_timeout(300)
        w2 = page.locator('.doc-sheet').bounding_box()['width']
        check('zoom 50 %', abs(w2 - 397) < 6, str(w2))
        check('pagination unchanged when zoomed', page.locator('.pv-page').count() == n)
        shot(page, 'writer_zoom50')
        page.locator('[data-zoom]').select_option('1')
        page.locator('[data-x="ruler"]').click()
        page.wait_for_timeout(200)
        check('ruler visible with indent markers', page.locator('.w-ruler .rl-first').is_visible() and page.locator('.w-ruler .rl-ticks i').count() > 50)
        # ruler: drag the left indent marker → paragraph indent
        wcaret(page, 'Absatz 1:', False)
        page.wait_for_timeout(100)
        k = page.locator('.rl-left').bounding_box()
        page.mouse.move(k['x'] + 6, k['y'] + 4)
        page.mouse.down()
        page.mouse.move(k['x'] + 46, k['y'] + 4, steps=4)
        page.mouse.up()
        page.wait_for_timeout(200)
        ml = page.evaluate("() => [...document.querySelectorAll('.doc-page p')].find(p => p.textContent.startsWith('Absatz 1:')).style.marginLeft")
        check('ruler drag sets left indent', ml.endswith('pt') and 25 < float(ml[:-2]) < 35, ml)
        # go to page
        page.evaluate(f"{W}.gotoPage(3)")
        page.wait_for_timeout(200)
        check('go to page 3 moves the caret there', page.evaluate(f"{W}.pv.pageOf(getSelection().anchorNode)") == 3)
        # web view
        page.locator('[data-x="view-web"]').click()
        page.wait_for_timeout(300)
        check('web view: no page backgrounds, no pagination marks', page.locator('.pv-page').count() == 0 and page.locator('.doc-page [data-pg]').count() == 0)
        page.locator('[data-x="view-page"]').click()
        page.wait_for_timeout(400)
        page.locator('[data-x="ruler"]').click()
        shot(page, 'writer_pages_light')
        # PDF page count agrees roughly
        model_pages = page.evaluate("async () => { const api = document.querySelector('.writer')._writerApi; const { modelToPdf } = await import('./js/office/docpdf.js'); const bytes = await modelToPdf(api.model(), { page: api.page(), hf: api.hf() }); const { loadPdfLib } = await import('./js/pdflib.js'); const L = await loadPdfLib(); return (await L.PDFDocument.load(bytes)).getPageCount(); }")
        check('PDF export page count close to page view', abs(model_pages - page.locator('.pv-page').count()) <= 1, f'{model_pages} vs {page.locator(".pv-page").count()}')
        leave(page)
    step('writer page layout', layout)

    def styles_fonts():
        new_doc(page)
        set_html(page, '<p>Mein Titel</p><p>Einleitung zum Thema</p><p>Normaler Text mit Wort fett und mehr.</p><p>Zweiter Absatz</p><p>klein und gross</p>')
        wcaret(page, 'Mein Titel')
        page.locator('.w-sty[data-sty="title"]').click()
        page.wait_for_timeout(100)
        check('style gallery: Titel', page.locator('.doc-page p.title', has_text='Mein Titel').count() == 1)
        wcaret(page, 'Einleitung')
        page.keyboard.press('Control+Shift+Digit1')
        page.wait_for_timeout(100)
        check('Ctrl+Shift+1 → Überschrift 1', page.locator('.doc-page h1', has_text='Einleitung').count() == 1)
        check('style select follows caret', page.locator('[data-block]').input_value() == 'h1')
        page.locator('[data-block]').select_option('blockquote')
        page.wait_for_timeout(100)
        check('style select: Zitat', page.locator('.doc-page p.quote', has_text='Einleitung').count() == 1)
        page.keyboard.press('Control+Shift+Digit2')
        page.wait_for_timeout(100)
        check('quote → Überschrift 2', page.locator('.doc-page h2', has_text='Einleitung').count() == 1 and page.locator('.doc-page p.quote').count() == 0)
        # modify style: heading 2 → Georgia 18pt red
        page.evaluate(f"{W}.setStyles('h2', {{ font: 'Georgia', size: 18, color: '#C01010' }})")
        page.wait_for_timeout(200)
        fs = page.evaluate("getComputedStyle(document.querySelector('.doc-page h2')).fontSize")
        ff = page.evaluate("getComputedStyle(document.querySelector('.doc-page h2')).fontFamily")
        check('modified style applies to all h2', fs == '24px' and 'Georgia' in ff, f'{fs} {ff}')
        # font family + size on selection
        wsel(page, 'Wort fett')
        page.locator('[data-font]').select_option('Courier New')
        page.wait_for_timeout(100)
        wsel(page, 'Wort fett')
        page.locator('[data-size]').select_option('14')
        page.wait_for_timeout(100)
        sp = page.evaluate("() => { const t = [...document.querySelectorAll('.doc-page span')].filter(s => s.textContent === 'Wort fett').pop(); return t ? getComputedStyle(t).fontFamily + '|' + getComputedStyle(t).fontSize : null; }")
        check('font family + size applied to selection', sp and 'Courier' in sp and sp.endswith('18.6667px'), str(sp))
        wsel(page, 'Wort fett')
        page.keyboard.press('Control+b')
        page.locator('[data-cmd="superscript"]').click()
        page.wait_for_timeout(100)
        check('bold + superscript', page.locator('.doc-page b, .doc-page strong').filter(has_text='Wort fett').count() >= 1 and page.locator('.doc-page sup', has_text='Wort fett').count() == 1)
        page.locator('[data-cmd="superscript"]').click()
        # change case
        wsel(page, 'klein und gross')
        page.locator('[data-x="case"]').click()
        menu(page, 'GROSSBUCHSTABEN')
        check('change case → upper', page.locator('.doc-page p', has_text='KLEIN UND GROSS').count() == 1)
        wsel(page, 'KLEIN UND GROSS')
        page.locator('[data-x="case"]').click()
        menu(page, 'Jedes Wort Groß')
        check('change case → title case', page.locator('.doc-page p', has_text='Klein Und Gross').count() == 1)
        # highlight
        wsel(page, 'Zweiter')
        page.locator('[data-color="back"]').click()
        menu(page, 'Markieren: Mint')
        check('highlight colour', page.evaluate("() => [...document.querySelectorAll('.doc-page span')].some(s => s.textContent === 'Zweiter' && /167, 243, 208/.test(s.style.backgroundColor))"))
        # format painter: copy from "Wort fett" (Courier, bold) to "Zweiter Absatz"
        wcaret(page, 'Wort fett', False)
        page.evaluate("() => { const s = getSelection(); s.modify('move', 'forward', 'character'); }")
        page.locator('[data-x="painter"]').click()
        wsel(page, 'Absatz', 0)
        page.locator('.doc-page').dispatch_event('mouseup')
        page.wait_for_timeout(200)
        painted = page.evaluate("() => { const r = getSelection().getRangeAt(0); const el = r.startContainer.nodeType === 1 ? r.startContainer : r.startContainer.parentElement; const t = [...document.querySelectorAll('.doc-page span, .doc-page b, .doc-page strong')].filter(s => s.textContent === 'Absatz').pop(); return t ? getComputedStyle(t).fontFamily + '|' + getComputedStyle(t).fontWeight : null; }")
        check('format painter transfers font + weight', painted and 'Courier' in painted and painted.endswith('700'), str(painted))
        # clear formatting
        wsel(page, 'Absatz', 0)
        page.locator('[data-x="clear"]').click()
        page.wait_for_timeout(100)
        cleared = page.evaluate("() => { const p = [...document.querySelectorAll('.doc-page p')].find(p => p.textContent.includes('Zweiter')); const t = [...p.querySelectorAll('*')].map(x => x.tagName + ':' + (x.getAttribute('style') || '')); return t; }")
        check('clear formatting removes character formats', not any('Courier' in t or 'font-weight' in t for t in cleared), str(cleared))
        # paragraph formatting via API (dialog values) → docx
        wcaret(page, 'Normaler Text')
        page.evaluate(f"{W}.F.applyPara([document.querySelector('.doc-page h2').nextElementSibling], {{ line: 1.5, before: 12, after: 18, indent: 36, firstLine: 18, align: 'justify', keepNext: true }})")
        hid = save_doc(page)
        path = save_blob(page, hid, 'styles.docx')
        d = Document(path)
        by = {p.text: p for p in d.paragraphs}
        check('docx: Title style', by.get('Mein Titel') is not None and by['Mein Titel'].style.name == 'Title', str({k: v.style.name for k, v in by.items()}))
        check('docx: Heading 2 style', by.get('Einleitung zum Thema') is not None and by['Einleitung zum Thema'].style.name == 'Heading 2')
        h2 = [x for x in d.styles if x.style_id == 'Heading2'][0]
        check('docx: modified Heading 2 definition (Georgia 18 pt red)', h2.font.name == 'Georgia' and h2.font.size == Pt(18) and str(h2.font.color.rgb) == 'C01010', f'{h2.font.name} {h2.font.size} {h2.font.color.rgb}')
        pf = [p for p in d.paragraphs if p.text.startswith('Normaler Text')][0].paragraph_format
        check('docx: paragraph spacing/indent/line/keep', pf.space_before == Pt(12) and pf.space_after == Pt(18) and pf.left_indent == Pt(36) and pf.first_line_indent == Pt(18) and abs((pf.line_spacing or 0) - 1.5) < 0.05 and pf.keep_with_next and int(pf.alignment) == 3, f'{pf.space_before} {pf.space_after} {pf.left_indent} {pf.first_line_indent} {pf.line_spacing} {pf.keep_with_next} {pf.alignment}')
        runs = [r for p in d.paragraphs for r in p.runs if 'Wort fett' in r.text]
        check('docx: run font Courier New 14 pt bold', runs and runs[0].font.name == 'Courier New' and runs[0].font.size == Pt(14) and runs[0].bold, str([(r.font.name, r.font.size, r.bold) for r in runs]))
        # reopen: styles + paragraph formats survive
        page.reload()
        page.wait_for_selector('.doc-page h2', timeout=15000)
        page.wait_for_timeout(500)
        ff2 = page.evaluate("getComputedStyle(document.querySelector('.doc-page h2')).fontFamily")
        st = page.evaluate("() => { const p = [...document.querySelectorAll('.doc-page p')].find(p => p.textContent.startsWith('Normaler Text')); return [p.style.lineHeight, p.style.marginLeft, p.style.textIndent, p.hasAttribute('data-keep-next')]; }")
        check('reopen: style definition from styles.xml', 'Georgia' in ff2, ff2)
        check('reopen: paragraph formats', st[1] == '36pt' and st[2] == '18pt' and st[3] and abs(float(st[0]) - 1.7) < 0.05, str(st))
        check('reopen: title style', page.locator('.doc-page p.title').count() == 1)
        leave(page)
    step('writer styles + fonts + paragraph', styles_fonts)

    def lists():
        new_doc(page)
        set_html(page, '<p>Eins</p><p>Zwei</p><p>Drei</p><p>Trenner</p><p>Alpha</p><p>Beta</p><p>Punkt A</p><p>Punkt B</p>')
        wsel(page, 'Eins')
        page.evaluate("() => { const ps = [...document.querySelectorAll('.doc-page p')]; const r = document.createRange(); r.setStart(ps[0].firstChild, 0); r.setEnd(ps[2].firstChild, 4); const s = getSelection(); s.removeAllRanges(); s.addRange(r); }")
        page.locator('[data-x="multilevel"]').click()
        menu(page, 'Mehrstufig: 1. / 1.1. / 1.1.1.')
        check('multilevel list created', page.locator('.doc-page ol[data-ml="legal"] > li').count() == 3)
        wcaret(page, 'Zwei', False)
        page.keyboard.press('Tab')
        page.wait_for_timeout(100)
        check('Tab indents list item', page.locator('.doc-page ol[data-ml="legal"] ol li', has_text='Zwei').count() == 1)
        marker = page.evaluate("() => { const li = [...document.querySelectorAll('.doc-page li')].find(l => l.textContent === 'Zwei'); return getComputedStyle(li, '::marker').content; }")
        check('legal numbering 1.1. shown', '1.1' in marker or 'counters' in marker, marker)
        page.evaluate("() => { const ps = [...document.querySelectorAll('.doc-page p')]; const a = ps.find(p => p.textContent === 'Alpha'), b = ps.find(p => p.textContent === 'Beta'); const r = document.createRange(); r.setStart(a.firstChild, 0); r.setEnd(b.firstChild, 4); const s = getSelection(); s.removeAllRanges(); s.addRange(r); }")
        page.locator('[data-x="multilevel"]').click()
        menu(page, 'Nummerierung: a. b. c.')
        check('lettered list', page.locator('.doc-page ol[type="a"] li').count() == 2)
        page.evaluate("() => { const ps = [...document.querySelectorAll('.doc-page p')]; const a = ps.find(p => p.textContent === 'Punkt A'), b = ps.find(p => p.textContent === 'Punkt B'); const r = document.createRange(); r.setStart(a.firstChild, 0); r.setEnd(b.firstChild, 7); const s = getSelection(); s.removeAllRanges(); s.addRange(r); }")
        page.locator('[data-x="multilevel"]').click()
        menu(page, 'Aufzählung: – Strich')
        check('dash bullets', page.locator('.doc-page ul[data-bullet="dash"] li').count() == 2)
        wcaret(page, 'Beta')
        page.locator('[data-x="multilevel"]').click()
        menu(page, 'Nummerierungswert festlegen')
        page.locator('dialog input[name=v]').fill('5')
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(200)
        check('restart numbering at 5 splits list', page.locator('.doc-page ol[start="5"] li', has_text='Beta').count() == 1)
        hid = save_doc(page)
        path = save_blob(page, hid, 'lists.docx')
        xml = docx_xml(path)
        numx = docx_xml(path, 'word/numbering.xml')
        check('docx: list levels (ilvl 1 for nested item)', 'w:ilvl w:val="1"' in xml)
        check('docx: legal multilevel format %1.%2.', '%1.%2.' in numx)
        check('docx: lowerLetter numbering', 'lowerLetter' in numx)
        check('docx: dash bullet', '–' in numx)
        check('docx: start override 5', 'w:start w:val="5"' in numx or 'w:startOverride w:val="5"' in numx)
        leave(page)
    step('writer lists', lists)

    def tables():
        new_doc(page)
        set_html(page, '<p>Vor der Tabelle</p>')
        wcaret(page, 'Vor der Tabelle')
        page.locator('[data-table]').click()
        page.locator('dialog [name=rows]').fill('3')
        page.locator('dialog [name=cols]').fill('3')
        page.locator('dialog [name=style]').select_option('striped')
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(300)
        check('table inserted with style + header row', page.locator('.doc-page table.ts-striped tr[data-hdr]').count() == 1)
        page.evaluate("() => { const tds = document.querySelectorAll('.doc-page td'); tds.forEach((td, i) => { td.innerHTML = '<p>Z' + i + '</p>'; }); }")
        wcaret(page, 'Z4')
        page.wait_for_timeout(150)
        check('contextual table tab appears', page.locator('.rb-tab[data-rb="table"]').is_visible())
        rb(page, 'table')
        page.locator('[data-t="row-below"]').click()
        page.locator('[data-t="col-right"]').click()
        page.wait_for_timeout(100)
        check('row below + column right', page.locator('.doc-page tr').count() == 4 and page.locator('.doc-page tr').first.locator('td').count() == 4)
        # select Z0..Z1 (two cells) and merge
        page.evaluate("() => { const tds = [...document.querySelectorAll('.doc-page td')]; const a = tds.find(t => t.textContent === 'Z0'), b = tds.find(t => t.textContent === 'Z1'); const r = document.createRange(); r.setStart(a.querySelector('p').firstChild, 0); r.setEnd(b.querySelector('p').firstChild, 2); const s = getSelection(); s.removeAllRanges(); s.addRange(r); document.dispatchEvent(new Event('selectionchange')); }")
        page.wait_for_timeout(200)
        check('cell range selection marked', page.locator('.doc-page td.cell-sel').count() == 2)
        page.locator('[data-t="merge"]').click()
        page.wait_for_timeout(150)
        check('merge cells → colspan 2', page.locator('.doc-page td[colspan="2"]').count() == 1)
        wcaret(page, 'Z7')
        page.locator('[data-t="split"]').click()
        page.locator('dialog [name=c]').fill('2')
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(150)
        check('split cell into 2 columns', page.evaluate("[...document.querySelectorAll('.doc-page tr')].find(r => r.textContent.includes('Z7')).cells.length") == 5, page.evaluate("document.querySelector('.doc-page table').outerHTML"))
        wcaret(page, 'Z8')
        page.locator('[data-t="shade"]').click()
        menu(page, 'Vanille')
        check('cell shading', page.evaluate("() => [...document.querySelectorAll('.doc-page td')].find(t => t.textContent === 'Z8').style.backgroundColor") == 'rgb(255, 244, 204)')
        page.locator('[data-t="borders"]').click()
        menu(page, 'Dicke Rahmenlinie unten')
        # column width drag
        tds = page.locator('.doc-page tr').nth(1).locator('td')
        bb = tds.nth(0).bounding_box()
        page.mouse.move(bb['x'] + bb['width'] - 1, bb['y'] + bb['height'] / 2)
        page.mouse.down()
        page.mouse.move(bb['x'] + bb['width'] + 60, bb['y'] + bb['height'] / 2, steps=5)
        page.mouse.up()
        page.wait_for_timeout(200)
        cols = page.evaluate("() => [...document.querySelectorAll('.doc-page col')].map(c => parseFloat(c.style.width))")
        check('column resize by dragging', len(cols) >= 4 and cols[0] > cols[2] + 5, str(cols))
        # undo / redo of a table operation
        nrows = page.locator('.doc-page tr').count()
        wcaret(page, 'Z8')
        page.locator('[data-t="row-above"]').click()
        page.wait_for_timeout(100)
        page.keyboard.press('Control+z')
        page.wait_for_timeout(150)
        u = page.locator('.doc-page tr').count()
        page.keyboard.press('Control+y')
        page.wait_for_timeout(150)
        check('undo/redo of table row insert', u == nrows and page.locator('.doc-page tr').count() == nrows + 1, f'{nrows} {u} {page.locator(".doc-page tr").count()}')
        page.keyboard.press('Control+z')
        hid = save_doc(page)
        path = save_blob(page, hid, 'table.docx')
        xml = docx_xml(path)
        check('docx: gridSpan from merge', 'w:gridSpan w:val="2"' in xml)
        check('docx: header row repeats (tblHeader)', '<w:tblHeader' in xml)
        check('docx: cell shading FFF4CC', 'FFF4CC' in xml.upper())
        check('docx: thick bottom border', 'w:sz="18"' in xml or 'w:sz="24"' in xml, '')
        d = Document(path)
        t = d.tables[0]
        check('docx: table opens with python-docx', len(t.rows) == 4 and t.cell(0, 0).text == 'Z0Z1' or 'Z0' in t.cell(0, 0).text, t.cell(0, 0).text)
        gridcols = re.findall(r'<w:gridCol w:w="(\d+)"', xml)
        check('docx: column widths follow the drag', len(gridcols) >= 4 and int(gridcols[0]) > int(gridcols[2]), str(gridcols))
        page.reload()
        page.wait_for_selector('.doc-page table', timeout=15000)
        check('reopen: merged cell + shading + header row', page.locator('.doc-page td[colspan]').count() >= 2 and page.locator('.doc-page tr[data-hdr]').count() == 1 and 'FFF4CC' in page.evaluate("document.querySelector('.doc-page table').outerHTML"), page.evaluate("document.querySelector('.doc-page table').outerHTML"))
        leave(page)
    step('writer tables', tables)

    def images_objects():
        new_doc(page)
        set_html(page, '<p>Text vor dem Bild. ' + 'Fülltext rund um das Bild. ' * 20 + '</p>')
        wcaret(page, 'Text vor dem Bild.')
        png = os.path.join(WORK, 'pic.png')
        from PIL import Image
        Image.new('RGB', (400, 300), (59, 55, 160)).save(png)
        with page.expect_file_chooser() as fc:
            page.locator('[data-image]').click()
        fc.value.set_files(png)
        page.wait_for_timeout(600)
        img = page.locator('.doc-page img').first
        img.click()
        page.wait_for_timeout(200)
        page.keyboard.press('Escape')
        check('image selected: handles + image tab', page.locator('.img-handles').is_visible() and page.locator('.rb-tab[data-rb="image"]').is_visible())
        w0 = img.bounding_box()['width']
        hb = page.locator('.img-handles [data-h="se"]').bounding_box()
        page.mouse.move(hb['x'] + 5, hb['y'] + 5)
        page.mouse.down()
        page.mouse.move(hb['x'] - 95, hb['y'] + 5, steps=5)
        page.mouse.up()
        page.wait_for_timeout(200)
        w1 = img.bounding_box()['width']
        check('resize image with handle', w1 < w0 - 60, f'{w0} → {w1}')
        rb(page, 'image')
        page.locator('[data-i="wrap"]').click()
        menu(page, 'Rechts, Text umfließt links')
        check('wrap right → float', page.evaluate("getComputedStyle(document.querySelector('.doc-page img')).float") == 'right')
        page.locator('[data-i="alt"]').click()
        page.wait_for_selector('dialog input[name=v]')
        page.wait_for_timeout(100)
        page.locator('dialog input[name=v]').fill('Blaues Rechteck')
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(150)
        check('alt text', page.locator('.doc-page img').first.get_attribute('alt') == 'Blaues Rechteck')
        # shape + text box
        wcaret(page, 'Text vor dem Bild.', False)
        rb(page, 'insert')
        page.locator('[data-x="ins-shape"]').click()
        menu(page, 'Ellipse')
        page.locator('[data-x="ins-textbox"]').click()
        page.wait_for_timeout(200)
        check('shape (SVG image) + text box inserted', page.locator('.doc-page img[alt="Ellipse"]').count() == 1 and page.locator('.doc-page .textbox').count() == 1)
        # columns
        set_html(page, page.evaluate("document.querySelector('.doc-page').innerHTML") + '<p>Spalte eins ' + 'Wort ' * 80 + '</p><p>Spalte zwei ' + 'Text ' * 80 + '</p>')
        page.evaluate("() => { const ps = [...document.querySelectorAll('.doc-page > p')]; const a = ps.find(p => p.textContent.startsWith('Spalte eins')), b = ps.find(p => p.textContent.startsWith('Spalte zwei')); const r = document.createRange(); r.setStart(a.firstChild, 0); r.setEnd(b.firstChild, 5); const s = getSelection(); s.removeAllRanges(); s.addRange(r); }")
        rb(page, 'layout')
        page.locator('[data-x="columns"]').click()
        menu(page, 'Zwei Spalten')
        check('two-column section', page.locator('.doc-page .sect[data-cols="2"] p').count() == 2 and page.evaluate("getComputedStyle(document.querySelector('.doc-page .sect')).columnCount") == '2')
        shot(page, 'writer_objects_light')
        hid = save_doc(page)
        path = save_blob(page, hid, 'objects.docx')
        xml = docx_xml(path)
        check('docx: floating image (anchor, wrap square)', '<wp:anchor' in xml and 'wrapSquare' in xml)
        check('docx: alt text', 'Blaues Rechteck' in xml)
        check('docx: two-column section', 'w:num="2"' in xml)
        d = Document(path)
        check('docx with objects opens in python-docx', len(d.paragraphs) > 2)
        leave(page)
    step('writer images, shapes, text boxes, columns', images_objects)

    def track():
        new_doc(page)
        set_html(page, '<p>Der schnelle braune Fuchs springt.</p><p>Zweite Zeile bleibt.</p>')
        rb(page, 'review')
        page.locator('[data-x="track"]').click()
        page.wait_for_timeout(100)
        wcaret(page, 'schnelle ', True)
        page.keyboard.type('sehr ')
        wsel(page, 'braune ')
        page.keyboard.press('Delete')
        wcaret(page, 'springt', True)
        page.keyboard.press('Backspace')
        page.keyboard.press('Backspace')
        page.wait_for_timeout(200)
        ins = page.locator('.doc-page ins.tc').all_inner_texts()
        dels = page.locator('.doc-page del.tc').all_inner_texts()
        check('tracked insertion', 'sehr ' in ins, str(ins))
        check('tracked deletions (selection + backspace)', 'braune ' in dels and any(x.endswith('gt') for x in dels), str(dels))
        check('deleted text still in DOM, author set', page.locator('.doc-page del.tc[data-author][data-date]').count() >= 2)
        page.locator('[data-x="changes"]').click()
        page.wait_for_timeout(200)
        check('changes panel lists changes', page.locator('.w-tc').count() >= 3, str(page.locator('.w-tc').count()))
        shot(page, 'writer_track_light')
        hid = save_doc(page)
        path = save_blob(page, hid, 'track.docx')
        xml = docx_xml(path)
        check('docx: w:ins + w:del + delText', '<w:ins ' in xml and '<w:del ' in xml and '<w:delText' in xml)
        check('docx: trackRevisions setting', 'trackRevisions' in docx_xml(path, 'word/settings.xml'))
        d = Document(path)
        check('docx with revisions opens', len(d.paragraphs) >= 2)
        page.reload()
        page.wait_for_selector('.doc-page', timeout=15000)
        page.wait_for_timeout(500)
        check('reopen: revisions shown', page.locator('.doc-page ins.tc').count() >= 1 and page.locator('.doc-page del.tc').count() >= 2)
        # reject the insertion, accept all the rest
        page.evaluate("() => { const ins = document.querySelector('.doc-page ins.tc'); const r = document.createRange(); r.selectNodeContents(ins); const s = getSelection(); s.removeAllRanges(); s.addRange(r); }")
        rb(page, 'review')
        page.locator('[data-x="reject"]').click()
        menu(page, 'Diese Änderung ablehnen')
        page.locator('[data-x="accept"]').click()
        menu(page, 'Alle Änderungen annehmen')
        txt = page.locator('.doc-page p').first.inner_text()
        check('reject insertion + accept deletions', txt == 'Der schnelle Fuchs sprin.' and page.locator('.doc-page ins.tc, .doc-page del.tc').count() == 0, txt)
        leave(page)
    step('writer track changes', track)

    def find_opts():
        new_doc(page)
        set_html(page, '<p>Haus Hausboot haus HAUS. Rechnung 2024-05 und 2025-11.</p>')
        page.keyboard.press('Control+f')
        page.locator('[data-fq]').fill('Haus')
        page.wait_for_timeout(200)
        a = page.locator('[data-fcount]').inner_text()
        page.locator('[data-fcase]').check()
        page.wait_for_timeout(150)
        b = page.locator('[data-fcount]').inner_text()
        page.locator('[data-fword]').check()
        page.wait_for_timeout(150)
        c = page.locator('[data-fcount]').inner_text()
        check('find: case-insensitive / match case / whole word', a.endswith('von 4') and b.endswith('von 2') and c.endswith('von 1'), f'{a} | {b} | {c}')
        page.locator('[data-fcase]').uncheck(); page.locator('[data-fword]').uncheck()
        page.locator('[data-fregex]').check()
        page.locator('[data-fq]').fill(r'(\d{4})-(\d{2})')
        page.locator('[data-fr]').fill('$2/$1')
        page.wait_for_timeout(150)
        page.locator('[data-freall]').click()
        page.wait_for_timeout(200)
        check('regex replace with groups', 'Rechnung 05/2024 und 11/2025.' in page.locator('.doc-page').inner_text(), page.locator('.doc-page').inner_text())
        page.locator('[data-fclose]').click()
        # word count dialog
        page.locator('[data-wordcount]').click()
        page.wait_for_timeout(200)
        tx = page.locator('dialog .w-stats').first.inner_text()
        check('word count dialog: words, characters, paragraphs, pages', 'Wörter' in tx and 'Zeichen (ohne Leerzeichen)' in tx and 'Absätze' in tx and 'Seiten' in tx, tx[:120])
        page.locator('dialog .btn-primary').click()
        page.keyboard.press('F1')
        page.wait_for_timeout(200)
        check('shortcut help sheet (F1)', page.locator('dialog .w-keys kbd').count() >= 15)
        page.locator('dialog .btn-primary').click()
        # keyboard: Ctrl+Enter page break, Ctrl+E center, Ctrl+K link dialog
        wcaret(page, 'HAUS.')
        page.keyboard.press('Control+e')
        page.wait_for_timeout(100)
        check('Ctrl+E centers paragraph', page.evaluate("getComputedStyle(document.querySelector('.doc-page p')).textAlign") == 'center')
        page.keyboard.press('Control+Enter')
        page.wait_for_timeout(200)
        check('Ctrl+Enter inserts a page break', page.locator('.doc-page hr[data-page-break]').count() == 1)
        page.keyboard.press('Control+z')
        page.wait_for_timeout(150)
        check('Ctrl+Z removes the page break again', page.locator('.doc-page hr[data-page-break]').count() == 0)
        leave(page)
    step('writer find options, word count, shortcuts', find_opts)

    def versions():
        new_doc(page)
        set_html(page, '<p>Version eins.</p>')
        hid = save_doc(page)
        set_html(page, '<p>Version zwei.</p>')
        page.evaluate("document.querySelector('.doc-page').dispatchEvent(new Event('input', {bubbles: true}))")
        save_doc(page)
        rb(page, 'view')
        page.locator('[data-x="versions"]').click()
        page.wait_for_timeout(300)
        n = page.locator('dialog .w-ver').count()
        check('version history lists saved versions', n >= 2, str(n))
        page.locator('dialog [data-restore="1"]').click()
        page.wait_for_timeout(200)
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(300)
        check('restore older version', 'Version eins.' in page.locator('.doc-page').inner_text())
        page.keyboard.press('Control+z')
        page.wait_for_timeout(200)
        check('restore is undoable', 'Version zwei.' in page.locator('.doc-page').inner_text())
        leave(page)
    step('writer versions', versions)

    def mailmerge():
        csv = os.path.join(WORK, 'empfaenger.csv')
        open(csv, 'w', encoding='utf-8').write('Anrede;Name;Ort\nFrau;Müller;Köln\nHerr;Schmidt;Berlin\nFrau;Öztürk;Hamburg\n')
        new_doc(page)
        set_html(page, '<p>Sehr geehrte(r) </p><p>wir grüßen nach .</p>')
        rb(page, 'mailings')
        page.locator('[data-x="mm-source"]').click()
        page.wait_for_timeout(300)
        with page.expect_file_chooser() as fc:
            page.locator('dialog button[value="device"]').click()
        fc.value.set_files(csv)
        page.wait_for_timeout(500)
        check('recipient list loaded', '1 / 3' in page.locator('[data-mmpos]').inner_text())
        wcaret(page, 'Sehr geehrte(r) ')
        page.locator('[data-x="mm-field"]').click()
        menu(page, 'Anrede')
        page.keyboard.type(' ')
        page.locator('[data-x="mm-field"]').click()
        menu(page, 'Name')
        wcaret(page, 'grüßen nach ')
        page.locator('[data-x="mm-field"]').click()
        menu(page, 'Ort')
        check('merge fields inserted', page.locator('.doc-page span.mf').count() == 3)
        page.locator('[data-x="mm-preview"]').click()
        page.wait_for_timeout(200)
        page.locator('[data-x="mm-next"]').click()
        page.wait_for_timeout(200)
        check('preview record 2', 'Herr Schmidt' in page.locator('.doc-page').inner_text() and 'Berlin' in page.locator('.doc-page').inner_text())
        shot(page, 'writer_mailmerge_light')
        page.locator('[data-x="mm-finish"]').click()
        page.wait_for_timeout(200)
        page.locator('dialog input[value="docx"]').check()
        with page.expect_download(timeout=30000) as dl:
            page.locator('dialog .btn-primary').click()
        zp = os.path.join(WORK, 'serienbrief.zip')
        dl.value.save_as(zp)
        z = zipfile.ZipFile(zp)
        names = z.namelist()
        texts = []
        for nm in names:
            open(os.path.join(WORK, 'mm.docx'), 'wb').write(z.read(nm))
            texts.append('\n'.join(p.text for p in Document(os.path.join(WORK, 'mm.docx')).paragraphs))
        check('mail merge ZIP: one docx per record', len(names) == 3, str(names))
        check('mail merge values filled (umlauts)', any('Frau Öztürk' in t and 'Hamburg' in t for t in texts) and any('Herr Schmidt' in t for t in texts), str(texts)[:200])
        page.locator('[data-x="mm-finish"]').click()
        page.wait_for_timeout(200)
        page.locator('dialog .btn-primary').click()
        page.wait_for_selector('dialog .result-mini', timeout=40000)
        with page.expect_download(timeout=20000) as d2:
            page.locator('dialog button[value="dl"]').click()
        pp = os.path.join(WORK, 'serienbrief.pdf')
        d2.value.save_as(pp)
        txt = subprocess.run(['pdftotext', pp, '-'], capture_output=True, text=True).stdout
        npages = int(re.search(r'Pages:\s+(\d+)', subprocess.run(['pdfinfo', pp], capture_output=True, text=True).stdout).group(1))
        check('mail merge PDF: one page per record', npages == 3 and 'Frau Müller' in txt and 'Köln' in txt, f'{npages} {txt[:80]}')
        leave(page)
    step('writer mail merge', mailmerge)

    def pdf_export():
        new_doc(page)
        rows = ''.join(f'<tr><td><p>Zeile {i}</p></td><td><p>{i * 3}</p></td></tr>' for i in range(70))
        set_html(page, '<p class="title">PDF-Test</p><p>Sichtbar <del class="tc" data-author="A" data-date="2025-01-01T00:00:00Z">GELÖSCHT</del><ins class="tc" data-author="A" data-date="2025-01-01T00:00:00Z">eingefügt</ins>.</p><div class="sect" data-cols="2" data-gap="18"><p>Spaltentext links ' + 'wort ' * 60 + '</p></div><table><tr data-hdr><td><p>Kopfzeile Tabelle</p></td><td><p>Wert</p></td></tr>' + rows + '</table><h2>Kapitel</h2><p>Ende.</p>')
        b64 = page.evaluate("async () => { const api = document.querySelector('.writer')._writerApi; const { modelToPdf } = await import('./js/office/docpdf.js'); const bytes = await modelToPdf(api.model(), { page: api.page(), hf: api.hf() }); let s = ''; const u = new Uint8Array(bytes); for (let i = 0; i < u.length; i += 0x8000) s += String.fromCharCode.apply(null, u.subarray(i, i + 0x8000)); return btoa(s); }")
        pp = os.path.join(WORK, 'export.pdf')
        open(pp, 'wb').write(base64.b64decode(b64))
        txt = subprocess.run(['pdftotext', pp, '-'], capture_output=True, text=True).stdout
        check('PDF: tracked deletion not printed, insertion printed', 'GELÖSCHT' not in txt and 'eingefügt' in txt)
        check('PDF: table header row repeated on following pages', txt.count('Kopfzeile Tabelle') >= 2, str(txt.count('Kopfzeile Tabelle')))
        check('PDF: column section + title style text present', 'Spaltentext links' in txt and 'PDF-Test' in txt)
        leave(page)
    step('writer PDF export details', pdf_export)

    def themes():
        new_doc(page)
        set_html(page, '<h1>Seitenlayout</h1><p>Normaler Absatz mit <b>fett</b>, <span style="background-color:#FDE047">markiert</span>.</p><ul><li>Eins</li><li>Zwei</li></ul><table class="ts-header"><tr data-hdr><td>Kopf</td><td>Wert</td></tr><tr><td>A</td><td>1</td></tr></table>')
        rb(page, 'view')
        page.locator('[data-x="ruler"]').click()
        shot(page, 'desktop_writer_light')
        set_theme(page, 'dark')
        shot(page, 'desktop_writer_dark')
        set_theme(page, 'light')
        page.locator('[data-x="ruler"]').click()
        leave(page)
    step('writer screenshots', themes)
    ctx.close()


def writer_phone(p):
    ctx = p.chromium.launch().new_context(viewport={'width': 390, 'height': 844}, device_scale_factor=2, is_mobile=True, has_touch=True, locale='de-DE', bypass_csp=True)
    page = ctx.new_page()
    attach_logging(page, 'phone')
    page.goto(URL)
    page.wait_for_selector('#view-start .hero')

    def ph():
        new_doc(page)
        set_html(page, '<h1>Handy</h1><p>' + 'Text auf dem Telefon. ' * 30 + '</p><table><tr><td>A</td><td>B</td></tr></table>')
        check('phone opens in web view', page.locator('.doc-sheet.pv').count() == 0)
        sw = page.evaluate('document.documentElement.scrollWidth')
        check('phone: no horizontal overflow', sw <= 390, str(sw))
        bb = page.locator('.doc-page').bounding_box()
        check('phone: page fits width', bb['width'] <= 390, str(bb['width']))
        shot(page, 'phone_writer_light')
        rb(page, 'view')
        page.locator('[data-x="view-page"]').tap()
        page.locator('[data-x="zoom-fit"]').tap()
        page.wait_for_timeout(400)
        bb = page.locator('.doc-sheet').bounding_box()
        check('phone: page layout fits width', bb['width'] <= 390 and page.locator('.pv-page').count() >= 1, str(bb['width']))
        shot(page, 'phone_writer_pages_light')
        set_theme(page, 'dark')
        shot(page, 'phone_writer_pages_dark')
        page.locator('[data-x="view-web"]').tap()
        shot(page, 'phone_writer_dark')
        set_theme(page, 'light')
        leave(page)
    step('writer phone', ph)
    ctx.close()


# =====================================================================
if __name__ == '__main__':
    with sync_playwright() as p:
        if ONLY in ('all', 'writer'):
            writer_tests(p)
        if ONLY in ('all', 'writer', 'phone'):
            writer_phone(p)
        if ONLY in ('all', 'sheet'):
            import sheet_tests2
            sheet_tests2.run(p, check, step, attach_logging, URL, OUT, WORK)
    print('\n==== CONSOLE errors/warnings ====')
    for c in console_problems[:40]:
        print(c)
    print(f'{len(console_problems)} console problems')
    passed = sum(1 for r in results if r[1])
    print(f'\n==== {passed}/{len(results)} passed ====')
    for r in results:
        if not r[1]:
            print('FAILED:', r[0], r[2])
