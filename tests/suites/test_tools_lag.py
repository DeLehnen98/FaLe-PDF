"""Tools must not freeze after files were added (phone, CPU ×4): main-thread stalls stay short.
- Seiten verwalten / Seiten extrahieren with a 30-page scanned PDF: thumbnails are rendered one after another
- Word → PDF and Text → PDF: drawing pauses, and the document index tokenizes long texts in pieces
- Bild-Werkzeuge and Passfoto with a 12 MP photo: preview made from the file off the main thread
- Zusammenfügen on a 360 px phone: rows fit, long names readable (two lines), buttons in reach
Server: node tests/suites/sec/csp_server.mjs 9202 <repo>"""
import base64, io, os, re, sys, traceback
import numpy as np
from PIL import Image, ImageDraw, ImageFilter
from playwright.sync_api import sync_playwright

URL = os.environ.get('APP_URL', 'http://127.0.0.1:9202/')
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'out_tools_lag')
os.makedirs(OUT, exist_ok=True)
LIMIT = 700  # ms, worst single stall at CPU ×4 (≈ a mid-range phone)
results, problems = [], []
INIT = "try { localStorage.setItem('fale-onboarded', '1'); localStorage.setItem('fale-install-hint', '1'); } catch (e) {}"
WATCH = """window.__d = []; (function () { let last = performance.now(); setInterval(() => { const n = performance.now(); const g = n - last - 50; if (g > 80) window.__d.push([Math.round(n), Math.round(g)]); last = n; }, 50); })();"""
IMPORT = """async ([b64, name, mime]) => { const bin = atob(b64); const u = new Uint8Array(bin.length); for (let i = 0; i < bin.length; i++) u[i] = bin.charCodeAt(i);
  const { importFiles } = await import('./js/library.js'); const [d] = await importFiles([new File([u], name, { type: mime })], { source: 'Gerät', any: true }); return d.id; }"""


def check(name, cond, info=''):
    results.append((name, bool(cond), str(info)))
    print(('PASS ' if cond else 'FAIL ') + name + (f'  [{str(info)[:400]}]' if info and not cond else ''), flush=True)


def step(name, fn):
    try:
        fn()
    except Exception as e:
        results.append((name + ' (exception)', False, str(e)[:300]))
        print('FAIL ' + name + ' EXC ' + str(e)[:600], flush=True)
        traceback.print_exc()


def files():
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    import docx
    rng = np.random.default_rng(3)
    f = {}
    f['scan'] = os.path.join(OUT, 'scan30.pdf')
    if not os.path.exists(f['scan']):
        c = canvas.Canvas(f['scan'], pagesize=A4)
        for p in range(30):
            a = (rng.random((2339, 1654)) * 25 + 215).astype('uint8')
            im = Image.fromarray(a, 'L').filter(ImageFilter.GaussianBlur(1.5)).convert('RGB')
            d = ImageDraw.Draw(im)
            for y in range(150, 2200, 60):
                d.text((150, y), 'Seite %d Zeile %d Mietvertrag Testtext' % (p + 1, y), fill=(0, 0, 0))
            b = io.BytesIO(); im.save(b, 'JPEG', quality=70); b.seek(0)
            c.drawImage(ImageReader(b), 0, 0, width=A4[0], height=A4[1]); c.showPage()
        c.save()
    f['photo'] = os.path.join(OUT, 'foto.jpg')
    if not os.path.exists(f['photo']):
        a = (rng.random((3024, 4032, 3)) * 60 + 150).astype('uint8')
        im = Image.fromarray(a).filter(ImageFilter.GaussianBlur(1.2))
        im.save(f['photo'], quality=90)
    f['docx'] = os.path.join(OUT, 'bericht.docx')
    D = docx.Document()
    for i in range(40):
        D.add_heading('Kapitel %d' % i, 1)
        for j in range(8):
            D.add_paragraph('Dies ist ein Absatz mit etwas Text für den Test. ' * 6)
    D.save(f['docx'])
    f['txt'] = os.path.join(OUT, 'notiz.txt')
    open(f['txt'], 'w', encoding='utf-8').write('Zeile mit Text äöü Kündigungsfrist drei Monate\n' * 3000)
    return f


def ctx(browser, w=412):
    c = browser.new_context(viewport={'width': w, 'height': 900}, locale='de-DE', device_scale_factor=2.6, is_mobile=True, has_touch=True)
    c.add_init_script(INIT + WATCH)
    pg = c.new_page()
    pg.on('pageerror', lambda e: problems.append(str(e)[:200]))
    pg.goto(URL + '#/start')
    pg.wait_for_selector('#view-start')
    pg.wait_for_timeout(1200)
    return c, pg


def worst(pg, t0):
    d = [x for x in pg.evaluate('window.__d') if x[0] > t0]
    return max([x[1] for x in d] or [0])


def run(p):
    F = files()
    browser = p.chromium.launch()

    def thumbs():
        for tid in ('seiten', 'extrahieren'):
            c, pg = ctx(browser)
            did = pg.evaluate(IMPORT, [base64.b64encode(open(F['scan'], 'rb').read()).decode(), 'scan30.pdf', 'application/pdf'])
            pg.wait_for_timeout(2500)
            c.new_cdp_session(pg).send('Emulation.setCPUThrottlingRate', {'rate': 4})
            t0 = pg.evaluate('performance.now()')
            pg.evaluate(f"location.hash = '#/werkzeug/{tid}?doc={did}'")
            pg.wait_for_timeout(10000)
            w = worst(pg, t0)
            n = pg.evaluate("document.querySelectorAll('#tool-screen canvas').length")
            check(f'{tid}: 30-page scan – thumbnails appear ({n}), no freeze (worst stall {w} ms < {LIMIT})', w < LIMIT and n >= 4, w)
            c.close()
    step('page thumbnails', thumbs)

    def convert():
        for tid, key in (('word-pdf', 'docx'), ('text-pdf', 'txt')):
            c, pg = ctx(browser)
            pg.evaluate(f"location.hash = '#/werkzeug/{tid}'")
            pg.wait_for_selector('#tool-screen [data-dev]', timeout=15000)
            pg.wait_for_timeout(1500)
            c.new_cdp_session(pg).send('Emulation.setCPUThrottlingRate', {'rate': 4})
            with pg.expect_file_chooser() as fc:
                pg.locator('#tool-screen [data-dev]').click()
            t0 = pg.evaluate('performance.now()')
            fc.value.set_files([F[key]])
            pg.wait_for_selector('#tool-screen [data-result] .result-card', timeout=90000)
            pg.wait_for_timeout(6000)  # the document index runs afterwards
            w = worst(pg, t0)
            check(f'{tid}: file added → PDF made, no freeze (worst stall {w} ms < {LIMIT})', w < LIMIT, w)
            c.close()
    step('office conversion', convert)

    def photos():
        for tid in ('bildwerkzeuge', 'passfoto'):
            c, pg = ctx(browser)
            pg.evaluate(f"location.hash = '#/werkzeug/{tid}'")
            pg.wait_for_timeout(2500)
            c.new_cdp_session(pg).send('Emulation.setCPUThrottlingRate', {'rate': 4})
            btn = pg.locator('#tool-screen button:visible').filter(has_text=re.compile(r'wählen', re.I)).first
            with pg.expect_file_chooser() as fc:
                btn.click()
            t0 = pg.evaluate('performance.now()')
            fc.value.set_files([F['photo']])
            pg.wait_for_timeout(9000)
            w = worst(pg, t0)
            shown = pg.evaluate("[...document.querySelectorAll('#tool-screen canvas')].some(c => c.offsetWidth > 100)")
            check(f'{tid}: 12 MP photo opens ({shown}), no long freeze (worst stall {w} ms < 1000)', shown and w < 1000, w)
            c.close()
    step('photos', photos)

    def merge_phone():
        c, pg = ctx(browser, 360)
        b64 = base64.b64encode(open(F['scan'], 'rb').read()).decode()
        for n in ('Mietvertrag mit einem langen Namen 2026.pdf', 'b.pdf'):
            pg.evaluate(IMPORT, [b64, n, 'application/pdf'])
        pg.evaluate("location.hash = '#/werkzeug/zusammenfuegen'")
        pg.wait_for_timeout(1500)
        pg.locator('#tool-screen button', has_text='Dokumente hinzufügen').click()
        pg.wait_for_selector('dialog[open]')
        items = pg.locator('dialog[open]').locator('.doc, button.doc-open, input[type=checkbox]')
        for i in range(min(items.count(), 2)):
            items.nth(i).click()
        pg.locator('dialog[open] .dlg-actions .btn-primary').first.click()
        pg.wait_for_selector('.order-item [data-remove]')
        pg.wait_for_timeout(800)
        wide = pg.evaluate("""() => { const W = document.documentElement.clientWidth; return [...document.querySelectorAll('#tool-screen *')].filter(e => { const r = e.getBoundingClientRect(); return r.width && r.right > W + 1; }).length; }""")
        check('Zusammenfügen 360 px: nothing sticks out on the right (buttons reachable)', wide == 0, wide)
        name_h = pg.evaluate("Math.max(...[...document.querySelectorAll('.order-item .doc-name')].map(e => e.getBoundingClientRect().width))")
        check('Zusammenfügen 360 px: names get the row width (not „Mie…“)', name_h > 180, name_h)
        sizes = pg.evaluate("[...document.querySelectorAll('.order-item .icon-btn')].map(b => Math.min(b.offsetWidth, b.offsetHeight))")
        check('Zusammenfügen 360 px: buttons at least 40 px', sizes and min(sizes) >= 40, sizes)
        before = pg.locator('.order-item .doc-name').all_inner_texts()
        pg.locator('.order-item [data-down]').first.click()
        pg.wait_for_timeout(300)
        after = pg.locator('.order-item .doc-name').all_inner_texts()
        check('Zusammenfügen: order changes with the arrow', after == before[::-1] and len(after) == 2, (before, after))
        pg.screenshot(path=os.path.join(OUT, 'merge_360.png'))
        c.close()
    step('merge on a phone', merge_phone)
    browser.close()


if __name__ == '__main__':
    with sync_playwright() as p:
        run(p)
    check('no page errors', not problems, problems[:5])
    bad = [r for r in results if not r[1]]
    print(f'\n{len(results) - len(bad)}/{len(results)} passed')
    for r in bad:
        print('  FAIL', r[0], r[2][:400])
    sys.exit(1 if bad else 0)
