"""pdf-to-excel (agent conv): invoice + 3-page table → XLSX with typed values, joined tables (repeated header
dropped), one sheet per page, CSV, whole-page mode / fallback, table preview."""
import csv, datetime, io
import openpyxl
from wtlib import Suite
from conv_testlib import shots, check_radio, back_to_files

P = {'de': '/werkzeuge/pdf-in-excel/', 'en': '/en/tools/pdf-to-excel/'}

with Suite('conv-excel') as s:
    for v in s.variants():
        page = s.open(v, P[v.lang])
        sh = v.device == 'phone' and v.lang == 'de'
        if sh: shots(s, page, 'pdf-to-excel', 'initial')
        s.pick(page, 'conv_tables.pdf'); s.wait_ready(page, 1)
        if sh: shots(s, page, 'pdf-to-excel', 'configured', top='.wt-opts')
        s.check_page(page)
        out = s.run_and_download(page)
        if sh: shots(s, page, 'pdf-to-excel', 'result', top='.wt-result', dark=True)
        s.check_page(page)
        if not out: continue
        wb = openpyxl.load_workbook(out)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        s.expect(rows[0][:5] == ('Pos.', 'Artikel', 'Menge', 'Einzelpreis', 'Gesamt'), f'{v}: invoice header {rows[0]}')
        s.expect(rows[1][:5] == (1, 'Schreibtischlampe', 2, 24.9, 49.8), f'{v}: invoice row typed {rows[1]}')
        s.expect(ws['D2'].number_format.startswith('#,##0.00'), f'{v}: currency format {ws["D2"].number_format}')
        s.expect(ws['A1'].font.b, f'{v}: header bold')
        dates = [r[0] for r in rows if isinstance(r[0], datetime.datetime)]
        s.expect(len(dates) == 70, f'{v}: 70 measurement rows joined over pages, got {len(dates)}')
        s.expect(sum(1 for r in rows if r[0] == 'Datum') == 1, f'{v}: repeated header rows dropped')
        pct = [r for r in rows if isinstance(r[0], datetime.datetime)][0]
        s.expect(pct[4] == 0 and isinstance(pct[2], (int, float)), f'{v}: % and numbers typed {pct}')
        s.expect(page.locator('.cv-table').count() >= 2, f'{v}: table preview')
        page.close()

    ctx = s.context('desktop', 'de')
    page = s.open(ctx, P['de'])
    s.pick(page, 'conv_tables.pdf'); s.wait_ready(page, 1)
    check_radio(page, 'wt-opt-sheets', 'pages')
    out = s.run_and_download(page)
    if out:
        wb = openpyxl.load_workbook(out)
        s.expect(wb.sheetnames == ['Seite 1', 'Seite 2', 'Seite 3'], f'sheets per page: {wb.sheetnames}')
    back_to_files(page)
    check_radio(page, 'wt-opt-sheets', 'one')
    check_radio(page, 'wt-opt-format', 'csv')
    out = s.run_and_download(page)
    if out:
        raw = out.read_bytes()
        s.expect(raw.startswith(b'\xef\xbb\xbf'), 'CSV has BOM')
        rows = list(csv.reader(io.StringIO(raw.decode('utf-8-sig')), delimiter=';'))
        s.expect(rows[1][:5] == ['1', 'Schreibtischlampe', '2', '24,90 €', '49,80 €'], f'CSV row {rows[1]}')
    page.close()
    # text without tables → whole page fallback with note
    page = s.open(ctx, P['de'])
    s.pick(page, 'text3.pdf'); s.wait_ready(page, 1)
    out = s.run_and_download(page)
    if out:
        s.expect('keine klaren Tabellen' in page.locator('.wt-result').inner_text(), 'fallback note')
        s.expect(openpyxl.load_workbook(out).active.max_row >= 3, 'fallback has rows')
    page.close()
    # scan → message
    page = s.open(ctx, P['de'])
    s.pick(page, 'conv_scan.pdf'); s.wait_ready(page, 1)
    s.primary(page).click()
    page.wait_for_selector('.wt-error:not([hidden])', timeout=30000)
    s.expect('gescannt' in page.locator('.wt-error').inner_text(), 'scan message')
    page.close()
