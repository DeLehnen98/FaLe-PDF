"""pdf-to-pptx (agent conv): one slide per page, page-shaped / 16:9 layouts, speaker notes, slide preview,
mixed page sizes, LibreOffice opens the file."""
import subprocess
from pptx import Presentation
from wtlib import Suite, OUT
from conv_testlib import shots, check_radio, toggle, back_to_files

P = {'de': '/werkzeuge/pdf-in-powerpoint/', 'en': '/en/tools/pdf-to-powerpoint/'}

with Suite('conv-pptx') as s:
    for v in s.variants():
        page = s.open(v, P[v.lang])
        sh = v.device == 'phone' and v.lang == 'de'
        if sh: shots(s, page, 'pdf-to-pptx', 'initial')
        s.pick(page, 'conv_report.pdf'); s.wait_ready(page, 1)
        if sh: shots(s, page, 'pdf-to-pptx', 'configured', top='.wt-opts')
        s.check_page(page)
        out = s.run_and_download(page)
        if sh: shots(s, page, 'pdf-to-pptx', 'result', top='.wt-result', dark=True)
        s.check_page(page)
        if not out: continue
        pr = Presentation(str(out))
        s.expect(len(pr.slides) == 2, f'{v}: 2 slides')
        ratio = pr.slide_width / pr.slide_height
        s.expect(abs(ratio - 595.28 / 841.89) < 0.01, f'{v}: page-shaped slides, ratio {ratio:.3f}')
        notes = pr.slides[0].notes_slide.notes_text_frame.text
        s.expect('Quartalsbericht Größenwachstum' in notes and 'Umsatz um 12 %' in notes, f'{v}: notes text')
        s.expect(page.locator('.cv-pages img').count() == 2, f'{v}: slide preview')
        page.close()

    ctx = s.context('desktop', 'de')
    page = s.open(ctx, P['de'])
    s.pick(page, 'conv_mixed.pdf'); s.wait_ready(page, 1)
    check_radio(page, 'wt-opt-layout', '16x9')
    toggle(page, 'wt-opt-notes', False)
    out = s.run_and_download(page)
    if out:
        pr = Presentation(str(out))
        s.expect(len(pr.slides) == 4 and abs(pr.slide_width / pr.slide_height - 16 / 9) < 0.01, '16:9 with 4 slides')
        s.expect(not pr.slides[0].has_notes_slide or not pr.slides[0].notes_slide.notes_text_frame.text, 'no notes when off')
        r = subprocess.run(['soffice', f'-env:UserInstallation=file://{OUT}/lo_profile', '--headless', '--convert-to', 'pdf', '--outdir', str(out.parent), str(out)], capture_output=True, timeout=180)
        s.expect(Suite.pdf_pages(out.with_suffix('.pdf')) == 4, 'LibreOffice renders 4 slides')
    s.check_page(page)
    page.close()
