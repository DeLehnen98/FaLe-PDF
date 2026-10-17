"""pdf-to-word (agent conv): editable DOCX (headings, bold/italic, lists, table, image, page break),
two columns, scanned page as image, exact mode, valid file (python-docx + LibreOffice renders it)."""
import os, subprocess, zipfile
from pathlib import Path
import docx
from wtlib import Suite, OUT, FIX
from conv_testlib import shots, check_radio, back_to_files

P = {'de': '/werkzeuge/pdf-in-word/', 'en': '/en/tools/pdf-to-word/'}
LO = f'-env:UserInstallation=file://{OUT}/lo_profile'


def lo_pages(path):
    subprocess.run(['soffice', LO, '--headless', '--convert-to', 'pdf', '--outdir', str(path.parent), str(path)], capture_output=True, timeout=180)
    pdf = path.with_suffix('.pdf')
    return Suite.pdf_pages(pdf) if pdf.exists() else -1


def line_rule_ok(path):
    """Every <w:spacing w:line=…> carries w:lineRule (else LibreOffice reads an exact height and clips pictures)."""
    import re
    z = zipfile.ZipFile(path)
    bad = []
    for n in ('word/document.xml', 'word/styles.xml'):
        if n in z.namelist():
            bad += [m for m in re.findall(r'<w:spacing\b[^>]*/>', z.read(n).decode()) if 'w:line=' in m and 'w:lineRule=' not in m]
    return bad


def ink(pdf, page=1):
    """Share of non-white pixels of one page of a PDF."""
    from PIL import Image
    prefix = OUT / f'_ink_{Path(pdf).stem}_{page}'
    subprocess.run(['pdftoppm', '-r', '20', '-f', str(page), '-l', str(page), '-png', '-singlefile', str(pdf), str(prefix)], check=True)
    px = Image.open(str(prefix) + '.png').convert('L').tobytes()
    return sum(1 for x in px if x < 235) / len(px)


with Suite('conv-word') as s:
    for v in s.variants():
        page = s.open(v, P[v.lang])
        sh = v.device == 'phone' and v.lang == 'de'
        if sh: shots(s, page, 'pdf-to-word', 'initial')
        s.pick(page, 'conv_report.pdf'); s.wait_ready(page, 1)
        if sh: shots(s, page, 'pdf-to-word', 'configured', top='.wt-opts')
        s.check_page(page)
        out = s.run_and_download(page)
        if sh: shots(s, page, 'pdf-to-word', 'result', top='.wt-result', dark=True)
        s.check_page(page)
        if not out: continue
        d = docx.Document(str(out))
        paras = [(p.style.name, p.text) for p in d.paragraphs]
        s.expect(('Heading 1', 'Quartalsbericht Größenwachstum') in paras, f'{v}: H1')
        s.expect(('Heading 2', 'Einleitung') in paras, f'{v}: H2')
        bold = [r.text for p in d.paragraphs for r in p.runs if r.bold]
        ital = [r.text for p in d.paragraphs for r in p.runs if r.italic]
        s.expect(any('fett gedruckte' in b for b in bold) and any('kursive' in i for i in ital), f'{v}: bold/italic runs')
        lists = [p.text for p in d.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]
        s.expect('Umsatz um 12 % gestiegen' in lists and 'Budget freigeben' in lists, f'{v}: list items with numbering: {lists}')
        s.expect(len(d.tables) == 1 and [c.text for c in d.tables[0].rows[1].cells] == ['Nord', '1.250,00 €', '800,50 €', '449,50 €'], f'{v}: table')
        s.expect(len(d.inline_shapes) == 1, f'{v}: one image, got {len(d.inline_shapes)}')
        xml = zipfile.ZipFile(out).read('word/document.xml').decode()
        s.expect('w:type="page"' in xml or 'pageBreakBefore' in xml, f'{v}: page break')
        s.expect(not line_rule_ok(out), f'{v}: line spacing written with w:lineRule: {line_rule_ok(out)[:3]}')
        s.expect(page.locator('.cv-doc').count() == 1 and 'Quartalsbericht' in page.locator('.cv-doc').inner_text(), f'{v}: structure preview')
        if v.lang == 'de' and v.device == 'desktop':
            s.expect(lo_pages(out) >= 2, 'LibreOffice opens and renders the DOCX')
        page.close()

    ctx = s.context('desktop', 'de')
    page = s.open(ctx, P['de'])
    s.pick(page, 'conv_columns.pdf'); s.wait_ready(page, 1)
    out = s.run_and_download(page)
    if out:
        txt = '\n'.join(p.text for p in docx.Document(str(out)).paragraphs)
        s.expect(txt.find('Links 3') < txt.find('Rechts 1') and 'Gemeindeblatt' not in txt, 'columns in order, header dropped')
    page.close()
    page = s.open(ctx, P['de'])
    s.pick(page, 'conv_scan.pdf'); s.wait_ready(page, 1)
    out = s.run_and_download(page)
    if out:
        s.expect(len(docx.Document(str(out)).inline_shapes) == 1, 'scan page inserted as image')
        s.expect('OCR' in page.locator('.wt-result').inner_text(), 'scan note mentions OCR')
    back_to_files(page)
    page.close()
    # exact mode on mixed page sizes
    page = s.open(ctx, P['de'])
    s.pick(page, 'conv_mixed.pdf'); s.wait_ready(page, 1)
    check_radio(page, 'wt-opt-mode', 'exact')
    out = s.run_and_download(page)
    if out:
        d = docx.Document(str(out))
        s.expect(len(d.sections) == 4 and len(d.inline_shapes) == 4, f'exact: 4 sections/pictures, got {len(d.sections)}')
        w = [round(sec.page_width.pt) for sec in d.sections]
        s.expect(w[0] == 595 and w[3] == 400, f'exact: page sizes {w}')
        s.expect(not line_rule_ok(out), f'exact: w:lineRule present: {line_rule_ok(out)[:3]}')
        s.expect(lo_pages(out) == 4, 'exact: LibreOffice renders 4 pages')
        got = [ink(out.with_suffix('.pdf'), n) for n in (1, 2, 3, 4)]
        want = [ink(FIX / 'conv_mixed.pdf', n) for n in (1, 2, 3, 4)]
        s.expect(all(g >= 0.7 * w for g, w in zip(got, want)), f'exact: LibreOffice shows the whole page pictures, not one-line strips ({got} vs {want})')
    s.check_page(page)
    page.close()
