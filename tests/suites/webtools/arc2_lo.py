"""arc2 helper: recalculate an XLSX in LibreOffice headless and read the computed values.

The FaLe XLSX writer may store cached formula results; to make LibreOffice really compute, every <v> of a formula
cell is removed first. LibreOffice then loads the file, calculates and saves it again as XLSX with fresh results,
which openpyxl reads (data_only=True)."""
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

PROFILE = Path(tempfile.gettempdir()) / 'arc2_lo_profile'


def strip_cached(src, dst):
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename.startswith('xl/worksheets/sheet'):
                xml = data.decode('utf-8')
                xml = re.sub(r'(<c [^>]*>)(<f[^>]*>.*?</f>|<f[^>]*/>)<v>[^<]*</v>', r'\1\2', xml)
                xml = re.sub(r'(<c [^>]*?) t="(str|e|b|n)"(?=[^>]*>\s*<f)', r'\1', xml)
                data = xml.encode('utf-8')
            zo.writestr(item, data)


def recalc(path, workdir):
    """→ path of the LibreOffice-recalculated copy (xlsx)."""
    workdir = Path(workdir)
    workdir.mkdir(parents=True, exist_ok=True)
    stripped = workdir / ('nocache_' + Path(path).name)
    strip_cached(path, stripped)
    outdir = workdir / 'lo'
    outdir.mkdir(exist_ok=True)
    r = subprocess.run(['soffice', f'-env:UserInstallation=file://{PROFILE}', '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', str(outdir), str(stripped)],
                       capture_output=True, text=True, timeout=180)
    out = outdir / stripped.name
    if not out.exists():
        raise RuntimeError('LibreOffice failed: ' + r.stderr[-400:] + r.stdout[-400:])
    return out


def values(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    return {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None} for ws in wb.worksheets}


def formulas(path):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    return {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=')} for ws in wb.worksheets}


if __name__ == '__main__':
    import sys
    p = recalc(sys.argv[1], tempfile.mkdtemp())
    v = values(p)
    for s, cells in v.items():
        print(s, list(cells.items())[:30])
    shutil.rmtree(PROFILE, ignore_errors=True)
