"""Fixtures for the conversion tools (agent `conv`). Writes fixtures/conv_*.

python3 conv_make_fixtures.py
"""
import io, os, subprocess, random
from pathlib import Path
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4, A5, landscape
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

FIX = Path(__file__).resolve().parent / 'fixtures'
FIX.mkdir(exist_ok=True)
DJ = '/usr/share/fonts/truetype/dejavu/'
pdfmetrics.registerFont(TTFont('DejaVu', DJ + 'DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('DejaVuB', DJ + 'DejaVuSans-Bold.ttf'))
random.seed(7)


def photo(w, h, seed=1):
    random.seed(seed)
    im = Image.new('RGB', (w, h))
    d = ImageDraw.Draw(im)
    for y in range(h):
        d.line([(0, y), (w, y)], fill=(int(40 + 180 * y / h), int(120 + 60 * ((seed * 37) % 3) / 3), int(220 - 150 * y / h)))
    for _ in range(25):
        x, y, r = random.randint(0, w), random.randint(0, h), random.randint(5, max(6, w // 6))
        d.ellipse([x - r, y - r, x + r, y + r], fill=(random.randint(0, 255), random.randint(0, 255), random.randint(0, 255)))
    return im


def jpeg_bytes(im, q=88, mode=None):
    b = io.BytesIO()
    (im.convert(mode) if mode else im).save(b, 'JPEG', quality=q)
    return b.getvalue()


# ---------------------------------------------------------------- report: headings, paragraphs, lists, table, images
def report():
    c = canvas.Canvas(str(FIX / 'conv_report.pdf'), pagesize=A4)
    c.setTitle('Quartalsbericht')
    W, H = A4
    y = H - 72

    def line(text, font='Helvetica', size=11, x=72, dy=None):
        nonlocal y
        c.setFont(font, size)
        c.drawString(x, y, text)
        y -= dy or size * 1.45

    line('Quartalsbericht Größenwachstum', 'Helvetica-Bold', 22, dy=34)
    line('Einleitung', 'Helvetica-Bold', 15, dy=22)
    para = ('Die Übersicht fasst die Ergebnisse des dritten Quartals zusammen. Wir erklären die wichtigsten '
            'Kennzahlen, die Entwicklung der Märkte und die nächsten Schritte für das kommende Jahr.')
    words = para.split()
    cur = ''
    for w in words:
        if c.stringWidth(cur + ' ' + w, 'Helvetica', 11) > W - 144:
            line(cur.strip())
            cur = ''
        cur += ' ' + w
    line(cur.strip())
    y -= 8
    c.setFont('Helvetica', 11)
    c.drawString(72, y, 'Dieser Satz enthält ')
    x = 72 + c.stringWidth('Dieser Satz enthält ', 'Helvetica', 11)
    c.setFont('Helvetica-Bold', 11); c.drawString(x, y, 'fett gedruckte'); x += c.stringWidth('fett gedruckte ', 'Helvetica-Bold', 11)
    c.setFont('Helvetica', 11); c.drawString(x, y, 'und '); x += c.stringWidth('und ', 'Helvetica', 11)
    c.setFont('Helvetica-Oblique', 11); c.drawString(x, y, 'kursive'); x += c.stringWidth('kursive ', 'Helvetica-Oblique', 11)
    c.setFont('Helvetica', 11); c.drawString(x, y, 'Wörter.')
    y -= 26
    line('Wichtige Punkte', 'Helvetica-Bold', 15, dy=22)
    for t in ['Umsatz um 12 % gestiegen', 'Neue Märkte in Österreich und der Schweiz', 'Kosten stabil geblieben']:
        line('•  ' + t, x=84)
    y -= 6
    for i, t in enumerate(['Budget freigeben', 'Team erweitern', 'Prüfung im Januar'], 1):
        line(f'{i}.  {t}', x=84)
    y -= 10
    line('Zahlen im Überblick', 'Helvetica-Bold', 15, dy=24)
    rows = [('Region', 'Umsatz', 'Kosten', 'Gewinn'), ('Nord', '1.250,00 €', '800,50 €', '449,50 €'),
            ('Süd', '2.300,00 €', '1.100,00 €', '1.200,00 €'), ('West', '980,75 €', '640,25 €', '340,50 €'),
            ('Summe', '4.530,75 €', '2.540,75 €', '1.990,00 €')]
    cols = [72, 200, 310, 420]
    for r, row in enumerate(rows):
        for k, cell in enumerate(row):
            c.setFont('Helvetica-Bold' if r == 0 else 'Helvetica', 10.5)
            c.drawString(cols[k], y, cell)
        c.setLineWidth(0.5); c.line(72, y - 5, 520, y - 5)
        y -= 18
    c.showPage()
    # page 2: images + more text
    y = H - 72
    line('Bilder aus dem Projekt', 'Helvetica-Bold', 15, dy=24)
    im = photo(800, 500, 3)
    c.drawImage(ImageReader(io.BytesIO(jpeg_bytes(im))), 72, y - 250, width=400, height=250)
    y -= 270
    line('Abbildung 1: Ein Foto als JPG eingebettet.', 'Helvetica-Oblique', 10)
    for k in range(6):
        line(f'Absatz {k + 1}: Hier steht ganz normaler Fließtext mit Umlauten wie äöü und ß, damit die Erkennung geprüft wird.', size=10.5)
    c.showPage()
    c.save()


# ---------------------------------------------------------------- images: jpeg rgb/gray/cmyk, png alpha, icons, duplicates
def images():
    c = canvas.Canvas(str(FIX / 'conv_images.pdf'), pagesize=A4)
    W, H = A4
    icon = Image.new('RGBA', (24, 24), (0, 0, 0, 0)); ImageDraw.Draw(icon).ellipse([2, 2, 22, 22], fill=(200, 30, 60, 255))
    logo_b = io.BytesIO(); icon.save(logo_b, 'PNG')
    alpha = Image.new('RGBA', (300, 200), (0, 0, 0, 0)); d = ImageDraw.Draw(alpha); d.rounded_rectangle([10, 10, 290, 190], 30, fill=(20, 120, 220, 200)); d.text((40, 90), 'ALPHA', fill=(255, 255, 255, 255))
    alpha_b = io.BytesIO(); alpha.save(alpha_b, 'PNG')
    gray = photo(400, 300, 5).convert('L')
    rgb = photo(640, 480, 9)
    for p in range(3):
        c.drawImage(ImageReader(io.BytesIO(logo_b.getvalue())), 40, H - 60, 24, 24, mask='auto')   # icon on every page
        c.setFont('Helvetica', 12); c.drawString(80, H - 52, f'Seite {p + 1}')
        if p == 0:
            c.drawImage(ImageReader(io.BytesIO(jpeg_bytes(rgb))), 72, H - 450, 320, 240)
            c.drawImage(ImageReader(io.BytesIO(alpha_b.getvalue())), 72, H - 700, 300, 200, mask='auto')
        if p == 1:
            c.drawImage(ImageReader(io.BytesIO(jpeg_bytes(gray))), 72, H - 400, 300, 225)
            c.drawImage(ImageReader(io.BytesIO(jpeg_bytes(rgb))), 300, H - 700, 160, 120)        # same photo again
        if p == 2:
            png = io.BytesIO(); photo(256, 256, 11).save(png, 'PNG')
            c.drawImage(ImageReader(io.BytesIO(png.getvalue())), 72, H - 400, 256, 256)
        c.showPage()
    c.save()
    # CMYK JPEG via pypdf? reportlab converts; use img2pdf-like raw embedding with pdf-lib later – keep a CMYK jpeg file
    jpeg_cmyk = jpeg_bytes(photo(200, 150, 4), mode='CMYK')
    (FIX / 'conv_cmyk.jpg').write_bytes(jpeg_cmyk)


# ---------------------------------------------------------------- tables: invoice and long list over pages
def tables():
    c = canvas.Canvas(str(FIX / 'conv_tables.pdf'), pagesize=A4)
    W, H = A4
    c.setFont('Helvetica-Bold', 16); c.drawString(72, H - 72, 'Rechnung Nr. 2024-117')
    c.setFont('Helvetica', 10.5); c.drawString(72, H - 92, 'Muster GmbH · Hauptstraße 5 · 66111 Saarbrücken')
    head = ('Pos.', 'Artikel', 'Menge', 'Einzelpreis', 'Gesamt')
    xs = [72, 110, 300, 370, 460]
    items = [('1', 'Schreibtischlampe', '2', '24,90 €', '49,80 €'), ('2', 'Bürostuhl Komfort', '1', '189,00 €', '189,00 €'),
             ('3', 'Druckerpapier A4', '10', '4,35 €', '43,50 €'), ('4', 'Kugelschreiber blau', '25', '0,89 €', '22,25 €'),
             ('5', 'Monitorhalterung', '1', '59,95 €', '59,95 €')]
    y = H - 140
    c.setFont('Helvetica-Bold', 10.5)
    for x, t in zip(xs, head):
        c.drawString(x, y, t)
    c.line(72, y - 4, 530, y - 4)
    y -= 18
    c.setFont('Helvetica', 10.5)
    for row in items:
        for k, (x, t) in enumerate(zip(xs, row)):
            if k >= 2:
                c.drawRightString(x + 60, y, t)
            else:
                c.drawString(x, y, t)
        y -= 16
    c.line(72, y + 10, 530, y + 10)
    c.setFont('Helvetica-Bold', 10.5); c.drawString(370, y - 4, 'Summe'); c.drawRightString(520, y - 4, '364,50 €')
    c.setFont('Helvetica', 10); c.drawString(72, y - 40, 'Zahlbar innerhalb von 14 Tagen ohne Abzug.')
    c.showPage()
    # page 2-3: long table with dates and numbers
    y = H - 72
    c.setFont('Helvetica-Bold', 14); c.drawString(72, y, 'Messwerte'); y -= 28
    hdr = ('Datum', 'Station', 'Temperatur', 'Niederschlag', 'Anteil')
    xs = [72, 160, 280, 370, 460]
    def header():
        nonlocal y
        c.setFont('Helvetica-Bold', 10)
        for x, t in zip(xs, hdr):
            c.drawString(x, y, t)
        y -= 16
    header()
    c.setFont('Helvetica', 10)
    for i in range(70):
        if y < 72:
            c.showPage(); y = H - 72; header(); c.setFont('Helvetica', 10)
        row = (f'{(i % 28) + 1:02d}.{(i // 28) + 3:02d}.2024', ['Nord', 'Süd', 'Ost', 'West'][i % 4], f'{12 + (i * 7) % 15},{i % 10}', f'{(i * 13) % 40},{(i * 3) % 10} mm', f'{(i * 17) % 100} %')
        for x, t in zip(xs, row):
            c.drawString(x, y, t)
        y -= 14
    c.showPage()
    c.save()


# ---------------------------------------------------------------- rotated / mixed sizes / long / non-latin
def misc():
    c = canvas.Canvas(str(FIX / 'conv_mixed.pdf'))
    for i, size in enumerate([A4, landscape(A4), A5, (400, 400)]):
        c.setPageSize(size)
        c.setFont('Helvetica-Bold', 20)
        c.drawString(40, size[1] - 60, f'Seite {i + 1} – {int(size[0])}×{int(size[1])}')
        c.setFont('Helvetica', 12)
        c.drawString(40, size[1] - 90, 'Unterschiedliche Seitengrößen und Ausrichtungen.')
        c.rect(20, 20, size[0] - 40, size[1] - 40)
        c.showPage()
    c.save()
    from pypdf import PdfReader, PdfWriter
    r = PdfReader(str(FIX / 'conv_mixed.pdf'))
    w = PdfWriter()
    for i, p in enumerate(r.pages):
        if i == 1:
            p.rotate(90)
        w.add_page(p)
    w.write(str(FIX / 'conv_mixed.pdf'))

    c = canvas.Canvas(str(FIX / 'conv_long.pdf'), pagesize=A4)
    for i in range(120):
        c.setFont('Helvetica-Bold', 18); c.drawString(72, 770, f'Kapitel {i + 1}')
        c.setFont('Helvetica', 11)
        for k in range(30):
            c.drawString(72, 740 - k * 20, f'Zeile {k + 1} auf Seite {i + 1}: Lorem ipsum dolor sit amet, äöü ß.')
        c.showPage()
    c.save()

    c = canvas.Canvas(str(FIX / 'conv_unicode.pdf'), pagesize=A4)
    c.setFont('DejaVuB', 18); c.drawString(72, 770, 'Многоязычный текст – Ελληνικά – Polski')
    c.setFont('DejaVu', 12)
    lines = ['Привет, мир! Это проверка извлечения текста.', 'Γειά σου Κόσμε – δοκιμή εξαγωγής κειμένου.',
             'Zażółć gęślą jaźń – Łódź, Kraków.', 'Türkçe: İstanbul, ğüşıöç. Čeština: příliš žluťoučký kůň.']
    for k, t in enumerate(lines):
        c.drawString(72, 730 - k * 24, t)
    c.showPage(); c.save()

    # scanned-like page (only an image, no text)
    im = Image.new('RGB', (1240, 1754), 'white'); d = ImageDraw.Draw(im)
    for k in range(30):
        d.text((100, 120 + k * 50), f'Gescannte Zeile {k + 1} – kein echter Text', fill=(30, 30, 30))
    c = canvas.Canvas(str(FIX / 'conv_scan.pdf'), pagesize=A4)
    c.drawImage(ImageReader(io.BytesIO(jpeg_bytes(im, 70))), 0, 0, *A4)
    c.showPage(); c.save()


# ---------------------------------------------------------------- office inputs
def office():
    import docx
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_BREAK
    d = docx.Document()
    d.core_properties.title = 'Testdokument'
    d.add_heading('Projektbericht Ölförderung & Süßwaren', 0)
    d.add_heading('1. Überblick', level=1)
    p = d.add_paragraph('Dies ist ein Absatz mit ')
    p.add_run('fettem').bold = True
    p.add_run(', ')
    p.add_run('kursivem').italic = True
    p.add_run(' und ')
    r = p.add_run('farbigem'); r.font.color.rgb = RGBColor(0xC0, 0x20, 0x40)
    p.add_run(' Text. Umlaute: ÄÖÜäöüß – „Anführungszeichen“ und das Euro-Zeichen 1.234,56 €.')
    for t in ['Erster Punkt', 'Zweiter Punkt', 'Dritter Punkt']:
        d.add_paragraph(t, style='List Bullet')
    for t in ['Schritt eins', 'Schritt zwei']:
        d.add_paragraph(t, style='List Number')
    d.add_heading('2. Zahlen', level=1)
    t = d.add_table(rows=4, cols=3)
    t.style = 'Table Grid'
    data = [('Monat', 'Umsatz', 'Änderung'), ('Januar', '12.500 €', '+3 %'), ('Februar', '13.100 €', '+4,8 %'), ('März', '11.900 €', '−9,2 %')]
    for i, row in enumerate(data):
        for k, v in enumerate(row):
            t.cell(i, k).text = v
    b = io.BytesIO(); photo(600, 400, 2).save(b, 'PNG'); b.seek(0)
    d.add_paragraph('Ein Bild:')
    d.add_picture(b, width=Inches(3))
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.add_heading('3. Zweite Seite', level=1)
    for k in range(12):
        d.add_paragraph(f'Absatz {k + 1}: Lorem ipsum dolor sit amet, consectetur adipiscing elit. Größere Mengen Text füllen die Seite, damit ein Umbruch entsteht. ' * 2)
    d.add_paragraph('Sonderzeichen außerhalb von WinAnsi: Привет, Ελληνικά, 漢字 ✓')
    d.save(str(FIX / 'conv_doc.docx'))

    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'Umsätze'
    ws.append(['Monat', 'Umsatz', 'Kosten', 'Gewinn', 'Datum'])
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDEEFF')
    import datetime
    for i, m in enumerate(['Januar', 'Februar', 'März', 'April', 'Mai', 'Juni', 'Juli', 'August', 'September', 'Oktober', 'November', 'Dezember'], 2):
        ws.append([m, 1000 + i * 137.5, 600 + i * 41.25, f'=B{i}-C{i}', datetime.date(2024, i - 1, 15)])
        ws[f'B{i}'].number_format = '#,##0.00 "€"'; ws[f'C{i}'].number_format = '#,##0.00 "€"'; ws[f'D{i}'].number_format = '#,##0.00 "€"'
        ws[f'E{i}'].number_format = 'DD.MM.YYYY'
    ws.append(['Summe', '=SUM(B2:B13)', '=SUM(C2:C13)', '=SUM(D2:D13)'])
    ws.column_dimensions['A'].width = 14; ws.column_dimensions['E'].width = 12
    ws2 = wb.create_sheet('Mitarbeiter')
    ws2.append(['Name', 'Abteilung', 'Stadt'])
    for k in range(60):
        ws2.append([f'Person {k + 1}', ['Vertrieb', 'Technik', 'Büro'][k % 3], ['Saarbrücken', 'Köln', 'München', 'Zürich'][k % 4]])
    wb.create_sheet('Leer')
    tmp = FIX / 'conv_sheet_raw.xlsx'
    wb.save(str(tmp))
    # let LibreOffice compute the formulas (cached values like a real Excel file)
    outdir = FIX / '_lo'
    outdir.mkdir(exist_ok=True)
    subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', str(outdir), str(tmp)], capture_output=True, timeout=120)
    if (outdir / 'conv_sheet_raw.xlsx').exists():
        (outdir / 'conv_sheet_raw.xlsx').rename(FIX / 'conv_sheet.xlsx')
    else:
        tmp.rename(FIX / 'conv_sheet.xlsx')
    if tmp.exists():
        tmp.unlink()

    from pptx import Presentation
    from pptx.util import Inches as PI, Pt as PP
    pr = Presentation()
    s = pr.slides.add_slide(pr.slide_layouts[0])
    s.shapes.title.text = 'Präsentation über Größen'
    s.placeholders[1].text = 'Untertitel mit Umlauten: äöü ß'
    s = pr.slides.add_slide(pr.slide_layouts[1])
    s.shapes.title.text = 'Agenda'
    tf = s.placeholders[1].text_frame
    tf.text = 'Begrüßung'
    for t in ['Zahlen & Fakten', 'Ausblick 2025']:
        tf.add_paragraph().text = t
    s = pr.slides.add_slide(pr.slide_layouts[5])
    s.shapes.title.text = 'Bild und Tabelle'
    b = io.BytesIO(); photo(400, 300, 6).save(b, 'PNG'); b.seek(0)
    s.shapes.add_picture(b, PI(0.5), PI(1.8), width=PI(4))
    tbl = s.shapes.add_table(3, 2, PI(5), PI(2), PI(4), PI(1.5)).table
    for i, row in enumerate([('Jahr', 'Wert'), ('2023', '42'), ('2024', '58')]):
        for k, v in enumerate(row):
            tbl.cell(i, k).text = v
    s.notes_slide.notes_text_frame.text = 'Sprechernotiz zur Folie 3'
    pr.save(str(FIX / 'conv_slides.pptx'))


def texts():
    (FIX / 'conv_text.txt').write_text(
        'Einkaufsliste für Samstag\n\nÄpfel, Birnen und Öl – 3,50 €\nSüßigkeiten für die Kinder\n\n'
        + 'Ein sehr langer Absatz, der umbrochen werden muss, weil er deutlich länger ist als eine Zeile auf einer A4-Seite. ' * 6
        + '\n\nUnicode: Привет мир · Ελληνικά · Zażółć gęślą jaźń · Emoji 🙂 · 漢字\n\tEingerückt mit Tab\n', encoding='utf-8')
    (FIX / 'conv_text.md').write_text(
        '# Überschrift eins\n\nEin Absatz mit **fettem** und *kursivem* Text sowie `Code`.\n\n## Liste\n\n- Erster Punkt\n- Zweiter Punkt\n  - Unterpunkt\n\n'
        '1. Eins\n2. Zwei\n\n## Tabelle\n\n| Name | Wert |\n| --- | --- |\n| Äpfel | 3 |\n| Birnen | 5 |\n\n> Ein Zitat.\n\n```\ncode block\n  eingerückt\n```\n', encoding='utf-8')
    (FIX / 'conv_text.csv').write_text('Name;Stadt;Betrag\nMüller;Köln;12,50\nÖzdemir;Saarbrücken;7,00\nŁukasz;Kraków;3,20\n', encoding='utf-8')
    (FIX / 'conv_latin1.txt').write_bytes('Grüße aus Saarbrücken – Größe 3½\n'.encode('cp1252'))


if __name__ == '__main__':
    report(); images(); tables(); misc(); office(); texts()
    print('ok', sorted(p.name for p in FIX.glob('conv_*')))


def cmyk_pdf():
    """Minimal PDF with a CMYK JPEG (DCTDecode, DeviceCMYK) – written by hand."""
    jpg = (FIX / 'conv_cmyk.jpg').read_bytes()
    w, h = 200, 150
    content = b'q 400 0 0 300 100 400 cm /Im1 Do Q'
    objs = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
            b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /XObject << /Im1 5 0 R >> >> /Contents 4 0 R >>',
            b'<< /Length %d >>\nstream\n' % len(content) + content + b'\nendstream',
            b'<< /Type /XObject /Subtype /Image /Width %d /Height %d /ColorSpace /DeviceCMYK /BitsPerComponent 8 /Decode [1 0 1 0 1 0 1 0] /Filter /DCTDecode /Length %d >>\nstream\n' % (w, h, len(jpg)) + jpg + b'\nendstream']
    out = bytearray(b'%PDF-1.4\n')
    offs = []
    for i, o in enumerate(objs, 1):
        offs.append(len(out)); out += b'%d 0 obj\n' % i + o + b'\nendobj\n'
    x = len(out)
    out += b'xref\n0 %d\n0000000000 65535 f \n' % (len(objs) + 1) + b''.join(b'%010d 00000 n \n' % o for o in offs)
    out += b'trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n' % (len(objs) + 1, x)
    (FIX / 'conv_cmyk.pdf').write_bytes(bytes(out))


if __name__ == '__main__':
    cmyk_pdf()


def columns_pdf():
    """Two-column newsletter: title across, columns, running header + page numbers on 4 pages."""
    import textwrap
    c = canvas.Canvas(str(FIX / 'conv_columns.pdf'), pagesize=A4)
    W, H = A4
    words = ('Die Gemeinde plant einen neuen Spielplatz am See. Bürgerinnen und Bürger können ihre Ideen bis '
             'Ende des Monats einreichen. Die Arbeiten sollen im Frühjahr beginnen und vor dem Sommer fertig sein. ').split()
    for p in range(4):
        c.setFont('Helvetica', 8); c.drawString(72, H - 40, 'Gemeindeblatt Saarland · Ausgabe 7')
        c.drawCentredString(W / 2, 30, f'Seite {p + 1} von 4')
        y0 = H - 80
        if p == 0:
            c.setFont('Helvetica-Bold', 20); c.drawString(72, y0, 'Neuigkeiten aus der Gemeinde am See'); y0 -= 40
        for col, x in enumerate((72, 316)):
            y = y0
            c.setFont('Helvetica-Bold', 13); c.drawString(x, y, f'Thema {p * 2 + col + 1}'); y -= 20
            c.setFont('Helvetica', 10)
            for para in range(3):
                txt = ' '.join(words[(para * 7 + col * 3) % len(words):] + words[:(para * 7 + col * 3) % len(words)])
                for ln in textwrap.wrap(f'{"Links" if col == 0 else "Rechts"} {para + 1}: ' + txt, 46):
                    c.drawString(x, y, ln); y -= 13
                y -= 9
        c.showPage()
    c.save()


if __name__ == '__main__':
    columns_pdf()
