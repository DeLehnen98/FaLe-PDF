"""Office fixtures for the Unicode tests (agent `uni`): DOCX / XLSX / PPTX with Polish, Greek, Cyrillic, Turkish, €."""
from pathlib import Path
FIX = Path(__file__).resolve().parent / 'fixtures'
UNI = 'Łukasz Kowalski-Müller Ελληνικά Привет Türkçe İğ €'

from docx import Document
d = Document()
d.add_heading('Łukasz Kowalski-Müller', level=1)
p = d.add_paragraph('Normal: ' + UNI + ' – ')
p.add_run('fett Ελληνικά').bold = True
p.add_run(' kursiv Привет').italic = True
t = d.add_table(rows=1, cols=2)
t.cell(0, 0).text = 'Türkçe İğ'
t.cell(0, 1).text = '100 €'
sec = d.sections[0]
sec.header.paragraphs[0].text = 'Kopf: Привет'
d.add_paragraph('Mono: Ελληνικά').runs[0].font.name = 'Courier New'
d.save(str(FIX / 'uni_doc.docx'))

import openpyxl
from openpyxl.styles import Font
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Łódź Привет'
ws['A1'] = 'Łukasz Kowalski-Müller'; ws['A1'].font = Font(bold=True)
ws['B1'] = 'Ελληνικά'
ws['A2'] = 'Привет'; ws['B2'] = 'Türkçe İğ'
ws['C2'] = '€ 12'
ws.column_dimensions['A'].width = 30
wb.save(str(FIX / 'uni_sheet.xlsx'))

from pptx import Presentation
from pptx.util import Inches
pr = Presentation()
s = pr.slides.add_slide(pr.slide_layouts[1])
s.shapes.title.text = 'Łukasz Kowalski-Müller'
s.placeholders[1].text = 'Ελληνικά Привет\nTürkçe İğ €'
pr.save(str(FIX / 'uni_slides.pptx'))
print('ok')
