"""Malicious test files for test_xss.py (written to secdata/).
Every payload sets window.__pwned = <id> if it ever executes; a DOM check also looks for injected markup."""
import os, io, zipfile, struct, zlib, json

SP = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(SP, 'secdata')
os.makedirs(OUT, exist_ok=True)


def P(n):  # HTML payload (as text)
    return f'<img src=x data-pwn="{n}" onerror="window.__pwned={n}">'


def X(s):  # XML-escape for attribute/text content
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('"', '&quot;')


W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def png_bytes(w=2, h=2):
    raw = b''.join(b'\x00' + b'\xff\x00\x00' * w for _ in range(h))
    def chunk(t, d):
        return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')


def write_zip(path, files, compress=zipfile.ZIP_DEFLATED):
    with zipfile.ZipFile(path, 'w', compress) as z:
        for name, data in files.items():
            z.writestr(name, data)


# ------------------------------------------------------------------ docx
doc_xml = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document {W}><w:body>
<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>{X(P(101))}</w:t></w:r></w:p>
<w:p><w:r><w:t xml:space="preserve">Text: {X(P(102))} &lt;script&gt;window.__pwned=103&lt;/script&gt;</w:t></w:r></w:p>
<w:p><w:hyperlink r:id="rId1"><w:r><w:t>JS-Link</w:t></w:r></w:hyperlink></w:p>
<w:p><w:hyperlink r:id="rId2"><w:r><w:t>Quote-Link</w:t></w:r></w:hyperlink></w:p>
<w:p><w:hyperlink r:id="rId6"><w:r><w:t>Data-Link</w:t></w:r></w:hyperlink></w:p>
<w:p><w:commentRangeStart w:id="1"/><w:r><w:t>Kommentiert</w:t></w:r><w:commentRangeEnd w:id="1"/></w:p>
<w:p><w:pPr><w:outlineLvl w:val="-7"/><w:jc w:val="center&quot; onclick=&quot;window.__pwned=104"/><w:shd w:fill="FF0000&quot; onmouseover=&quot;window.__pwned=105"/></w:pPr><w:r><w:rPr><w:color w:val="00FF00;background:url(https://evil.example/x)"/><w:rFonts w:ascii="x&quot; onmouseover=&quot;window.__pwned=106"/></w:rPr><w:t>Formatiert</w:t></w:r></w:p>
<w:p><w:r><w:drawing><wp:inline><wp:extent cx="95250" cy="95250"/><wp:docPr id="1" name="Bild" descr="{X('" onerror="window.__pwned=107" x="')}"/><a:graphic><a:graphicData><pic:pic><pic:blipFill><a:blip r:embed="rId4"/></pic:blipFill></pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>
<w:p><w:r><w:drawing><wp:inline><wp:extent cx="95250" cy="95250"/><wp:docPr id="2" name="Extern"/><a:graphic><a:graphicData><pic:pic><pic:blipFill><a:blip r:link="rId7"/></pic:blipFill></pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>
<w:sectPr/></w:body></w:document>'''
doc_rels = f'''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="{REL}/hyperlink" Target="javascript:window.__pwned=110" TargetMode="External"/>
<Relationship Id="rId2" Type="{REL}/hyperlink" Target="https://example.com/&quot;&gt;&lt;img src=x onerror=window.__pwned=111&gt;" TargetMode="External"/>
<Relationship Id="rId6" Type="{REL}/hyperlink" Target="data:text/html,&lt;script&gt;parent.__pwned=112&lt;/script&gt;" TargetMode="External"/>
<Relationship Id="rId3" Type="{REL}/comments" Target="comments.xml"/>
<Relationship Id="rId4" Type="{REL}/image" Target="media/bild.png"/>
<Relationship Id="rId7" Type="{REL}/image" Target="https://evil.example/track.png" TargetMode="External"/>
<Relationship Id="rId5" Type="{REL}/styles" Target="styles.xml"/>
</Relationships>'''
comments = f'''<?xml version="1.0" encoding="UTF-8"?>
<w:comments {W}><w:comment w:id="1" w:author="{X(P(120))}" w:date="&quot;&gt;&lt;img src=x onerror=window.__pwned=121&gt;"><w:p><w:r><w:t>{X(P(122))}</w:t></w:r></w:p></w:comment></w:comments>'''
styles = f'''<?xml version="1.0" encoding="UTF-8"?>
<w:styles {W}><w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/></w:style>
<w:style w:type="paragraph" w:styleId="{X(P(130))}"><w:name w:val="{X(P(131))}"/></w:style></w:styles>'''
ct = '''<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'''
pkg_rels = f'''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="{REL}/officeDocument" Target="word/document.xml"/></Relationships>'''
write_zip(os.path.join(OUT, 'xss.docx'), {'[Content_Types].xml': ct, '_rels/.rels': pkg_rels, 'word/document.xml': doc_xml,
                                          'word/_rels/document.xml.rels': doc_rels, 'word/comments.xml': comments, 'word/styles.xml': styles, 'word/media/bild.png': png_bytes()})

# ------------------------------------------------------------------ zip bombs (docx)
base = {'[Content_Types].xml': ct, '_rels/.rels': pkg_rels, 'word/_rels/document.xml.rels': '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>'}
# a) honest headers: 600 MiB of zeros declared (entry limit 512 MiB)
with zipfile.ZipFile(os.path.join(OUT, 'bomb_big.docx'), 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as z:
    for k, v in base.items():
        z.writestr(k, v)
    with z.open('word/document.xml', 'w', force_zip64=True) as f:
        blk = b'\0' * (8 * 1024 * 1024)
        for _ in range(75):
            f.write(blk)
# b) lying headers: 64 MiB inflate, but the zip claims 1 000 bytes
buf = io.BytesIO()
with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as z:
    for k, v in base.items():
        z.writestr(k, v)
    z.writestr('word/document.xml', b'<' + b' ' * (64 * 1024 * 1024))
data = bytearray(buf.getvalue())
real = 64 * 1024 * 1024 + 1
needle = struct.pack('<I', real)
n = 0
i = data.find(needle)
while i >= 0:
    data[i:i + 4] = struct.pack('<I', 1000)
    n += 1
    i = data.find(needle, i + 4)
assert n >= 2, n
open(os.path.join(OUT, 'bomb_lie.docx'), 'wb').write(bytes(data))

# ------------------------------------------------------------------ xlsx
SS = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
x_ct = '''<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>
<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>
<Override PartName="/xl/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/></Types>'''
x_rels = f'''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="{REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>'''
wb = f'''<?xml version="1.0" encoding="UTF-8"?>
<workbook {SS}><sheets><sheet name="{X('Blatt<b>' + P(201)[:20])}" sheetId="1" r:id="rId1"/></sheets>
<definedNames><definedName name="Name1" comment="{X(P(202))}">Blatt1!$A$1</definedName></definedNames></workbook>'''
wb_rels = f'''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="{REL}/worksheet" Target="worksheets/sheet1.xml"/>
<Relationship Id="rId2" Type="{REL}/styles" Target="styles.xml"/>
<Relationship Id="rId3" Type="{REL}/theme" Target="theme/theme1.xml"/>
<Relationship Id="rId4" Type="{REL}/sharedStrings" Target="sharedStrings.xml"/></Relationships>'''
sheet = f'''<?xml version="1.0" encoding="UTF-8"?>
<worksheet {SS}><sheetData>
<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" s="1"><v>42</v></c><c r="C1" s="2"><v>43</v></c><c r="D1" t="inlineStr"><is><t>{X(P(203))}</t></is></c><c r="E1" s="3"><v>44</v></c></row>
<row r="2"><c r="A2" t="str"><f>"&lt;img src=x onerror=window.__pwned=204&gt;"</f><v>{X(P(204))}</v></c></row>
</sheetData></worksheet>'''
sst = f'''<?xml version="1.0" encoding="UTF-8"?>
<sst {SS} count="1" uniqueCount="1"><si><t>{X(P(205))}</t></si></sst>'''
xstyles = f'''<?xml version="1.0" encoding="UTF-8"?>
<styleSheet {SS}>
<numFmts count="1"><numFmt numFmtId="164" formatCode="&quot;&lt;img src=x onerror=window.__pwned=206&gt;&quot;0"/></numFmts>
<fonts count="3"><font><sz val="11"/><name val="Calibri"/></font>
<font><sz val="11"/><name val="{X('x\'" data-q="1"><b data-pwn="207">F</b><i x="')}"/></font>
<font><sz val="11"/><color theme="4"/><name val="Calibri"/></font></fonts>
<fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills>
<borders count="1"><border/></borders>
<cellXfs count="4"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/><xf numFmtId="0" fontId="1" fillId="0" borderId="0" applyFont="1"/><xf numFmtId="0" fontId="2" fillId="0" borderId="0" applyFont="1"/><xf numFmtId="164" fontId="0" fillId="0" borderId="0" applyNumberFormat="1"/></cellXfs>
</styleSheet>'''
theme = f'''<?xml version="1.0" encoding="UTF-8"?>
<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="T"><a:themeElements><a:clrScheme name="S">
<a:dk1><a:srgbClr val="000000"/></a:dk1><a:lt1><a:srgbClr val="FFFFFF"/></a:lt1><a:dk2><a:srgbClr val="44546A"/></a:dk2><a:lt2><a:srgbClr val="E7E6E6"/></a:lt2>
<a:accent1><a:srgbClr val="{X('000000"><b data-pwn="208">T</b><i x="')}"/></a:accent1><a:accent2><a:srgbClr val="ED7D31"/></a:accent2>
</a:clrScheme></a:themeElements></a:theme>'''
write_zip(os.path.join(OUT, 'xss.xlsx'), {'[Content_Types].xml': x_ct, '_rels/.rels': x_rels, 'xl/workbook.xml': wb, 'xl/_rels/workbook.xml.rels': wb_rels,
                                          'xl/worksheets/sheet1.xml': sheet, 'xl/sharedStrings.xml': sst, 'xl/styles.xml': xstyles, 'xl/theme/theme1.xml': theme})

# ------------------------------------------------------------------ pptx (python-pptx base + raw injections)
from pptx import Presentation
from pptx.util import Inches
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = P(301)
s.placeholders[1].text = 'Untertitel ' + P(302)
tb = s.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1))
tb.text_frame.text = 'Box ' + P(303)
b = io.BytesIO(); prs.save(b)
zin = zipfile.ZipFile(io.BytesIO(b.getvalue()))
files = {n: zin.read(n) for n in zin.namelist()}
x = files['ppt/slides/slide1.xml'].decode('utf-8')
x = x.replace('type="ctrTitle"', 'type="' + X('ctrTitle"><b data-pwn="304">P</b><i x="') + '"', 1)
x = x.replace('<a:t>Box', '<a:rPr><a:latin typeface="' + X("x';}</style><img src=x onerror=window.__pwned=305>") + '"/><a:solidFill><a:srgbClr val="' + X('FF0000;background:url(https://evil.example/y)') + '"/></a:solidFill></a:rPr><a:t>Box', 1)
files['ppt/slides/slide1.xml'] = x.encode('utf-8')
write_zip(os.path.join(OUT, 'xss.pptx'), files)

# ------------------------------------------------------------------ markdown / text / csv
open(os.path.join(OUT, 'xss.md'), 'w').write(f'''# Titel {P(401)}

<script>window.__pwned=402</script>
<img src=x onerror="window.__pwned=403">

[js-link](javascript:window.__pwned=404) [data-link](data:text/html,<script>alert(1)</script>) [ok](https://example.com/"onmouseover="window.__pwned=405)

![bild](javascript:window.__pwned=406)

| a | {P(407)} |
|---|---|
| `{P(408)}` | **{P(409)}** |
''')
open(os.path.join(OUT, 'xss.csv'), 'w').write(f'Name,Wert\n"{P(501)}",1\n"=HYPERLINK(""javascript:window.__pwned=502"",""x"")",2\n')

# ------------------------------------------------------------------ images: dimension bomb (header 60000 x 60000, tiny data)
def big_png(w, h):
    def chunk(t, d):
        return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(b'\x00' * 64)) + chunk(b'IEND', b'')
open(os.path.join(OUT, 'huge.png'), 'wb').write(big_png(60000, 60000))
open(os.path.join(OUT, 'small.png'), 'wb').write(png_bytes(4, 4))
print('secdata ok:', sorted(os.listdir(OUT)))
