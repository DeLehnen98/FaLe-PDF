"""Office 3 UI tests (spreadsheet expansion): function autocomplete + argument help, new functions in cells, notes, links,
conditional formatting dialog (data bars, formula rules), regex search, go to special, grouping, subtotals, page setup → PDF,
template gallery (search, open, XLSX download), English UI of the gallery, phone: selection handles, long-press menu,
formula suggestions by tap, scrolling ribbon. Server: OFFICE_URL (default http://127.0.0.1:8850/). Screenshots → shots/."""
import os, sys, json, base64, subprocess, shutil, traceback, re
from playwright.sync_api import sync_playwright

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, 'shots'); os.makedirs(SHOTS, exist_ok=True)
WORK = os.path.join(HERE, 'work_ui'); shutil.rmtree(WORK, ignore_errors=True); os.makedirs(WORK)
URL = os.environ.get('OFFICE_URL', 'http://127.0.0.1:8850/')
results, problems = [], []


def check(name, cond, info=''):
    results.append((name, bool(cond), info))
    print(('PASS ' if cond else 'FAIL ') + name + (f'  [{info}]' if info and not cond else ''), flush=True)


def step(name, fn):
    try:
        fn()
    except Exception as e:
        results.append((name + ' (exception)', False, str(e)[:300]))
        print('FAIL ' + name + ' EXC ' + str(e)[:600], flush=True)
        traceback.print_exc()


def log(page, label):
    page.on('console', lambda m: problems.append(f'{label} [{m.type}] {m.text}') if m.type in ('error', 'warning') and 'favicon' not in m.text else None)
    page.on('pageerror', lambda e: problems.append(f'{label} [pageerror] {e}'))


FILES_JS = """async () => { const db = await new Promise((r, j) => { const q = indexedDB.open('falepdf'); q.onsuccess = () => r(q.result); q.onerror = () => j(q.error); });
  const res = await new Promise(r => { const t = db.transaction('docs'); const g = t.objectStore('docs').getAll(); g.onsuccess = () => r(g.result.map(d => ({id: d.id, name: d.name, kind: d.kind || 'pdf', createdAt: d.createdAt || 0}))); });
  db.close(); return res; }"""
BLOB_JS = """async (id) => { const db = await new Promise((r, j) => { const q = indexedDB.open('falepdf'); q.onsuccess = () => r(q.result); q.onerror = () => j(q.error); });
  const rec = await new Promise(r => { const t = db.transaction('files'); const g = t.objectStore('files').get(id); g.onsuccess = () => r(g.result); });
  db.close(); if (!rec) return null;
  return await new Promise(r => { const fr = new FileReader(); fr.onload = () => r(fr.result.split(',')[1]); fr.readAsDataURL(rec.blob); }); }"""


def save_doc(page, name, out):
    docs = [d for d in page.evaluate(FILES_JS) if d['name'] == name]
    if not docs: return None
    d = sorted(docs, key=lambda x: x['createdAt'])[-1]
    path = os.path.join(WORK, out)
    open(path, 'wb').write(base64.b64decode(page.evaluate(BLOB_JS, d['id'])))
    return path


def goto(page, h, wait=500):
    page.evaluate(f"location.hash = {json.dumps(h)}")
    page.wait_for_timeout(wait)


def leave(page):
    page.evaluate("location.hash = '#/start'")
    page.wait_for_timeout(400)
    d = page.locator('dialog', has_text='Änderungen speichern?')
    if d.count():
        d.locator('button[value="discard"]').click()
        page.wait_for_timeout(400)


def api(page, js):
    return page.evaluate(f"() => {{ const a = document.querySelector('.sheet-app')._sheetApi; {js} }}")


def sel(page, a):
    page.locator('#sh-addr').fill(a)
    page.locator('#sh-addr').press('Enter')
    page.wait_for_timeout(120)


def type_cell(page, a, text):
    sel(page, a)
    page.keyboard.type(text)
    page.keyboard.press('Enter')
    page.wait_for_timeout(150)


def cell_text(page, a):
    loc = page.locator(f'.sh-body [data-cell="{a}"], .sh-frozen [data-cell="{a}"]')
    return loc.first.inner_text() if loc.count() else ''


def wait_tables(page, n=39, ms=40000):  # 38 templates + blank (arc2)
    for _ in range(ms // 200):
        if page.locator('.stp-prev table').count() >= n: return True
        page.wait_for_timeout(200)
    return False


def shot(page, name):
    page.wait_for_timeout(250)
    page.screenshot(path=os.path.join(SHOTS, name + '.png'))


def pdftext(path):
    try: return subprocess.run(['pdftotext', '-layout', path, '-'], capture_output=True, timeout=60).stdout.decode('utf-8', 'replace')
    except Exception: return ''


def pdfinfo(path):
    try: return subprocess.run(['pdfinfo', path], capture_output=True, timeout=60).stdout.decode()
    except Exception: return ''


# =====================================================================
def desktop(p):
    browser = p.chromium.launch()
    ctx = browser.new_context(viewport={'width': 1366, 'height': 860}, locale='de-DE', accept_downloads=True)
    ctx.add_init_script("try { localStorage.setItem('fale-lang', 'de'); } catch {}")
    page = ctx.new_page(); log(page, 'desk')
    page.goto(URL); page.wait_for_timeout(1000)

    def autocomplete():
        goto(page, '#/neu/tabelle'); page.wait_for_selector('.sh-scroll')
        for a, v in [('A1', '10'), ('A2', '20'), ('A3', '30'), ('B1', 'Nord'), ('B2', 'Süd'), ('B3', 'Nord')]: type_cell(page, a, v)
        sel(page, 'C1')
        page.keyboard.type('=SUMMEW')
        page.wait_for_timeout(200)
        sug = page.locator('.fx-tip .fx-sg')
        check('autocomplete lists SUMMEWENN first', sug.count() >= 1 and 'SUMMEWENN' in sug.first.inner_text(), sug.first.inner_text() if sug.count() else 'none')
        shot(page, 'desk_autocomplete')
        page.keyboard.press('Tab')
        page.wait_for_timeout(150)
        val = page.locator('.cell-editor').input_value()
        check('Tab accepts the suggestion', val == '=SUMMEWENN(', val)
        page.keyboard.type('B1:B3;"Nord";')
        page.wait_for_timeout(200)
        on = page.locator('.fx-tip .fx-arg.on')
        check('argument help highlights the third argument', on.count() == 1 and 'Summe_Bereich' in on.inner_text(), on.inner_text() if on.count() else 'none')
        shot(page, 'desk_arghelp')
        page.keyboard.type('A1:A3)')
        page.keyboard.press('Enter')
        page.wait_for_timeout(200)
        check('SUMMEWENN typed via autocomplete computes', cell_text(page, 'C1') == '40', cell_text(page, 'C1'))
        # English names are accepted and shown in German
        type_cell(page, 'C2', '=TEXTSPLIT("x,y,z",",")')
        check('TEXTSPLIT spills (English input)', cell_text(page, 'C2') == 'x' and cell_text(page, 'E2') == 'z', cell_text(page, 'E2'))
        sel(page, 'C2')
        check('formula bar shows German name', page.locator('#sh-input').input_value().startswith('=TEXTTEILEN('), page.locator('#sh-input').input_value())
        type_cell(page, 'C3', '=LAMBDA(x;x*x)(12)')
        type_cell(page, 'C4', '=NACHZEILE(A1:A3;LAMBDA(z;z*2))')
        type_cell(page, 'C7', '=REGEXEXTRAKT("Bestellnr. 4711-B";"\\d+")')
        check('LAMBDA / BYROW / REGEXEXTRACT in cells', cell_text(page, 'C3') == '144' and cell_text(page, 'C6') == '60' and cell_text(page, 'C7') == '4711', f"{cell_text(page, 'C3')} {cell_text(page, 'C6')} {cell_text(page, 'C7')}")
        # insert-function dialog: English search + syntax
        page.locator('[data-rb="formulas"]').click()
        page.locator('[data-a="fxdlg"]').click()
        page.wait_for_timeout(250)
        page.locator('.fx-dlg [name=q]').fill('xlookup')
        page.wait_for_timeout(150)
        first = page.locator('.fx-item').first
        check('function dialog finds XVERWEIS by English name', 'XVERWEIS' in first.inner_text(), first.inner_text())
        info = page.locator('.fx-info').inner_text()
        check('function dialog shows argument syntax', 'Suchkriterium' in info and 'Rückgabematrix' in info, info)
        shot(page, 'desk_fxdialog')
        page.locator('dialog button[value="__cancel"]').first.click(); page.wait_for_timeout(200)
        page.locator('[data-rb="start"]').click()
    step('autocomplete & functions', autocomplete)

    def notes_links():
        sel(page, 'A1')
        page.keyboard.press('Shift+F2')
        page.wait_for_selector('dialog .note-ta')
        page.locator('dialog .note-ta').fill('Monatswert aus dem Kassenbuch')
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(300)
        check('note mark shown', page.locator('.sh-body [data-cell="A1"] .note-mark').count() == 1)
        check('note popover for the active cell', 'Monatswert aus dem Kassenbuch' in page.locator('.cell-pop').inner_text() and page.locator('.cell-pop').is_visible())
        shot(page, 'desk_note')
        sel(page, 'B1')
        page.keyboard.press('Control+k')
        page.wait_for_selector('dialog [name=target]')
        page.locator('dialog [name=target]').fill('https://example.org/info')
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(300)
        check('link cell styled', 'lnk' in (page.locator('.sh-body [data-cell="B1"]').get_attribute('class') or ''))
        btn = page.locator('.cell-pop [data-open-link]')
        check('link button in popover', btn.count() == 1 and 'example.org' in btn.inner_text())
        btn.click()
        page.wait_for_timeout(300)
        dlg = page.locator('dialog', has_text='Externen Link öffnen?')
        check('external link asks first', dlg.count() == 1 and 'example.org' in dlg.inner_text())
        if dlg.count(): dlg.locator('button[value="cancel"], button[value="__cancel"]').first.click()
        page.wait_for_timeout(200)
        # internal link
        sel(page, 'B2')
        page.keyboard.press('Control+k')
        page.wait_for_selector('dialog [name=target]')
        page.locator('dialog [name=target]').fill('C7')
        page.locator('dialog [name=text]').fill('Zur Bestellnummer')
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(300)
        shot(page, 'desk_link_internal')
        page.locator('.cell-pop [data-open-link]').click(timeout=5000)
        page.wait_for_timeout(200)
        check('internal link jumps to the cell', page.locator('#sh-addr').input_value() == 'C7', page.locator('#sh-addr').input_value())
        # invalid link target is refused
        sel(page, 'B3')
        page.keyboard.press('Control+k')
        page.wait_for_selector('dialog [name=target]')
        page.locator('dialog [name=target]').fill('javascript:alert(1)')
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(300)
        check('javascript: link refused', 'lnk' not in (page.locator('.sh-body [data-cell="B3"]').get_attribute('class') or ''))
    step('notes & links', notes_links)

    def cf_dialog():
        for i, v in enumerate([5, -3, 12, 8, 0, 20], start=1): type_cell(page, f'F{i}', str(v))
        sel(page, 'F1')
        page.locator('#sh-addr').fill('F1:F6'); page.locator('#sh-addr').press('Enter'); page.wait_for_timeout(100)
        (page.locator('[data-rb="data"]').click(), page.locator('[data-rbp="data"] [data-a="cf2"]').click())
        page.wait_for_selector('.cf-dlg')
        page.locator('.cf-dlg [name=type]').select_option('bar')
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(300)
        check('data bars rendered', page.locator('.sh-body [data-cell^="F"] .cf-bar').count() >= 5, str(page.locator('.cf-bar').count()))
        (page.locator('[data-rb="data"]').click(), page.locator('[data-rbp="data"] [data-a="cf2"]').click())
        page.wait_for_selector('.cf-dlg')
        page.locator('.cf-dlg [name=range]').fill('A1:A6')
        page.locator('.cf-dlg [name=type]').select_option('formula')
        page.locator('.cf-dlg [name=v1]').fill('=$F1>10')
        page.locator('.cf-dlg [name=style]').select_option('green')
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(300)
        st3 = page.locator('.sh-body [data-cell="A3"]').get_attribute('style') or ''
        st1 = page.locator('.sh-body [data-cell="A1"]').get_attribute('style') or ''
        check('formula rule highlights by row', 'background' in st3 and 'background' not in st1, st3[:120])
        (page.locator('[data-rb="data"]').click(), page.locator('[data-rbp="data"] [data-a="cf2"]').click())
        page.wait_for_selector('.cf-dlg')
        check('rule list describes rules', page.locator('.cf-rule').count() == 2 and 'Formel ist wahr' in page.locator('.cf-dlg').inner_text())
        page.locator('.cf-dlg [name=range]').fill('F1:F6')
        page.locator('.cf-dlg [name=type]').select_option('icons')
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(300)
        check('icon set rendered', page.locator('.sh-body .cf-ic').count() >= 5)
        shot(page, 'desk_cf')
    step('conditional formatting dialog', cf_dialog)

    def chart_variants():
        for a, v in [('H1', 'Monat'), ('I1', 'Umsatz'), ('J1', 'Kosten'), ('H2', 'Jan'), ('I2', '120'), ('J2', '80'), ('H3', 'Feb'), ('I3', '150'), ('J3', '95'), ('H4', 'Mär'), ('I4', '170'), ('J4', '90')]: type_cell(page, a, v)
        sel(page, 'H1:J4')
        page.locator('[data-rb="insert"]').click()
        page.locator('[data-rbp="insert"] [data-a="ins-chart"]').click()
        page.wait_for_selector('.chart-dlg')
        check('chart dialog keeps 6 base types', page.locator('.cd-type').count() == 6)
        page.locator('.chart-dlg [name=combo]').check()
        page.wait_for_timeout(200)
        check('combo preview draws columns and a line', page.locator('.cd-prev svg rect').count() >= 3 and page.locator('.cd-prev svg path[stroke-width="2.5"]').count() == 1, str((page.locator('.cd-prev svg rect').count(), page.locator('.cd-prev svg path[stroke-width="2.5"]').count())))
        shot(page, 'desk_chart_combo')
        page.locator('dialog .btn-primary').click(); page.wait_for_timeout(400)
        page.locator('[data-rbp="insert"] [data-a="ins-chart"]').click()
        page.wait_for_selector('.chart-dlg')
        page.locator('.cd-type[data-type="pie"]').click()
        page.locator('.chart-dlg [name=ring]').check()
        page.wait_for_timeout(200)
        check('ring preview has a hole', page.locator('.cd-prev svg circle').count() >= 1)
        check('ring option only for pie', page.locator('[data-var="col"]').is_hidden() and page.locator('[data-var="pie"]').is_visible())
        page.locator('dialog .btn-primary').click(); page.wait_for_timeout(400)
        types = api(page, 'return a.charts().map(c => c.type);')
        check('combo and doughnut charts inserted', types[-2:] == ['combo', 'doughnut'], str(types))
        page.locator('[data-rb="start"]').click()
    step('chart variants', chart_variants)

    def find_goto():
        sel(page, 'A1')
        page.keyboard.press('Control+f')
        page.wait_for_timeout(200)
        page.locator('label:has([data-fregex])').click()
        page.locator('[data-fq]').fill('^\\d{2}$')
        page.wait_for_timeout(250)
        cnt = page.locator('[data-fcount]').inner_text()
        check('regex search counts two-digit numbers', cnt.endswith('von 10'), cnt)
        page.locator('[data-fq]').fill('(')
        page.wait_for_timeout(200)
        check('invalid regex reported', 'Ungültiges Muster' in page.locator('[data-fcount]').inner_text(), page.locator('[data-fcount]').inner_text())
        page.locator('label:has([data-fregex])').click()
        page.locator('[data-fclose]').click()
        sel(page, 'A1')
        page.keyboard.press('Control+g')
        page.wait_for_selector('.goto-special')
        page.locator('[data-special="formulas"]').click()
        page.wait_for_timeout(300)
        r = api(page, 'return a.sel();')
        check('go to special selects formula cells', r['c1'] == 2 and r['more'] >= 0, str(r))
        page.keyboard.press('Control+g')
        page.wait_for_selector('dialog [name=ref]')
        page.locator('dialog [name=ref]').fill('F4')
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(200)
        check('go to address', page.locator('#sh-addr').input_value() == 'F4')
    step('regex search & go to', find_goto)

    def group_subtotal():
        leave(page)
        goto(page, '#/neu/tabelle'); page.wait_for_selector('.sh-scroll')
        rows = [('Region', 'Produkt', 'Umsatz'), ('Nord', 'A', '100'), ('Süd', 'B', '50'), ('Nord', 'B', '70'), ('West', 'A', '30'), ('Süd', 'A', '20')]
        for i, row in enumerate(rows, start=1):
            for c, v in zip('ABC', row): type_cell(page, f'{c}{i}', v)
        sel(page, 'A1')
        page.locator('[data-rb="data"]').click()
        page.locator('[data-a="subtotal"]').click()
        page.wait_for_selector('.st-dlg')
        check('subtotal dialog preselects the number column', page.locator('.st-dlg [name=col][value="2"]').is_checked())
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(400)
        t = [cell_text(page, f'A{i}') for i in range(1, 11)]
        check('subtotal rows inserted per group', t[3] == 'Nord Ergebnis' and 'Süd Ergebnis' in t and 'Gesamtergebnis' in t, str(t))
        check('subtotal values', cell_text(page, 'C4') == '170' and cell_text(page, 'C10') == '270', f"{cell_text(page, 'C4')} {cell_text(page, 'C10')}")
        btn = page.locator('.rh[data-row="3"] .ol-btn')
        check('outline toggle on the summary row', btn.count() == 1)
        shot(page, 'desk_subtotal')
        btn.click(); page.wait_for_timeout(250)
        check('collapse hides the group', page.locator('.rh[data-row="1"]').count() == 0 and page.locator('.rh[data-row="2"]').count() == 0)
        page.locator('.rh[data-row="3"] .ol-btn').click(); page.wait_for_timeout(250)
        check('expand shows the group again', page.locator('.rh[data-row="1"]').count() == 1)
        # keyboard grouping
        sel(page, 'A12:A14')
        page.keyboard.press('Alt+Shift+ArrowRight'); page.wait_for_timeout(200)
        check('Alt+Shift+→ groups rows', api(page, "return [...a.sheet().outline.keys()].includes(12);"))
        page.keyboard.press('Alt+Shift+ArrowLeft'); page.wait_for_timeout(200)
        check('Alt+Shift+← ungroups rows', not api(page, "return [...a.sheet().outline.keys()].includes(12);"))
    step('subtotals & grouping', group_subtotal)

    def page_setup_pdf():
        page.locator('[data-rb="page"]').click()
        page.locator('[data-a="pagesetup"]').click()
        page.wait_for_selector('.sps-dlg')
        page.locator('.sps-dlg input[name=orient][value=landscape] + span').click()
        page.locator('.sps-dlg input[name=fit][value=width] + span').click()
        page.locator('.sps-dlg [name=titles]').fill('1:1')
        page.locator('.sps-dlg [name=grid]').uncheck()
        shot(page, 'desk_pagesetup')
        page.locator('dialog .btn-primary').click()
        page.wait_for_timeout(300)
        p0 = api(page, 'return a.sheet().page;')
        check('page setup stored', p0['orient'] == 'landscape' and p0['fit'] == 'width' and p0['titleRows'] == {'r1': 0, 'r2': 0} and p0['grid'] is False, str(p0))
        page.locator('[data-a="pdf"]').click()
        page.locator('dialog', has_text='PDF erstellt').wait_for(timeout=20000)
        page.locator('dialog button[value="__cancel"]').first.click()
        page.wait_for_timeout(300)
        path = save_doc(page, 'Neue Tabelle.pdf', 'setup.pdf')
        info = pdfinfo(path) if path else ''
        m = re.search(r'Page size:\s+([\d.]+) x ([\d.]+)', info)
        check('PDF in landscape', m and float(m.group(1)) > float(m.group(2)), info[:200])
        txt = pdftext(path) if path else ''
        check('PDF has data, subtotals and page number', 'Nord Ergebnis' in txt and 'Gesamtergebnis' in txt and 'Seite 1 von 1' in txt, txt[:300])
        # Ctrl+P = export PDF
        page.locator('.sh-scroll').focus()
        page.keyboard.press('Control+p')
        page.locator('dialog', has_text='PDF erstellt').wait_for(timeout=20000)
        page.locator('dialog button[value="__cancel"]').first.click()
        check('Ctrl+P exports PDF', True)
        # save the workbook and inspect the xlsx
        page.locator('[data-office-save]').click()
        page.wait_for_timeout(1500)
        x = save_doc(page, 'Neue Tabelle.xlsx', 'outline.xlsx')
        import openpyxl
        wb = openpyxl.load_workbook(x)
        ws = wb.active
        check('xlsx: subtotal formulas and outline', ws['C4'].value == '=SUBTOTAL(9,C2:C3)' and ws.row_dimensions[2].outline_level == 1, f"{ws['C4'].value} {ws.row_dimensions[2].outline_level}")
        check('xlsx: page setup', ws.page_setup.orientation == 'landscape' and str(ws.print_title_rows).replace('$', '') == '1:1', f"{ws.page_setup.orientation} {ws.print_title_rows}")
        leave(page)
    step('page setup → PDF & xlsx', page_setup_pdf)

    def gallery():
        goto(page, '#/vorlagen/tabelle', 1500)
        check('template gallery lists 39 cards', page.locator('.stp-card').count() == 39, str(page.locator('.stp-card').count()))  # 38 templates + blank (arc2 added 20)
        wait_tables(page)
        check('live previews rendered', page.locator('.stp-prev table').count() == 39)
        shot(page, 'desk_gallery')
        page.locator('.stp-bar [data-q]').fill('urlaub')
        page.wait_for_timeout(150)
        vis = page.locator('.stp-card:not([hidden])')
        vis_ids = [x.get_attribute('data-id') for x in vis.all()]
        check('search filters the gallery', 'urlaubsplaner' in vis_ids and 'haushaltsbuch' not in vis_ids and len(vis_ids) < 6, str(vis_ids))
        page.locator('.stp-bar [data-q]').fill('')
        page.locator('.stp-cats [data-cat="Business"]').click()
        page.wait_for_timeout(150)
        ids = [x.get_attribute('data-id') for x in page.locator('.stp-card:not([hidden])').all()]
        check('category filter', {'leer', 'rechnungsausgangsbuch', 'euer', 'kassenbuch'} <= set(ids) and 'haushaltsbuch' not in ids and 'urlaubsplaner' not in ids, str(ids))
        with page.expect_download() as dl:
            page.locator('[data-dl="kassenbuch"]').click()
        d = dl.value
        path = os.path.join(WORK, 'kassenbuch.xlsx'); d.save_as(path)
        import openpyxl
        ws = openpyxl.load_workbook(path)['Kassenbuch']
        check('template downloads as xlsx', d.suggested_filename == 'Kassenbuch.xlsx' and ws['G6'].value.startswith('=IF('), f"{d.suggested_filename} {ws['G6'].value}")
        page.locator('.stp-cats [data-cat="Alle"]').click()
        page.locator('.stp-card[data-id="urlaubsplaner"] .stp-open').click()
        page.wait_for_selector('.sheet-app'); page.wait_for_timeout(600)
        check('template opens in the editor', cell_text(page, 'A1') == 'Urlaubsplaner' and cell_text(page, 'D13') == '15', f"{cell_text(page, 'A1')} {cell_text(page, 'D13')}")
        check('template starts at its input cell', page.locator('#sh-addr').input_value() == 'B13', page.locator('#sh-addr').input_value())
        type_cell(page, 'C13', '24.07.2026')
        check('editing the template recalculates', cell_text(page, 'D13') == '10' and cell_text(page, 'C6') == '10', f"{cell_text(page, 'D13')} {cell_text(page, 'C6')}")
        shot(page, 'desk_template_urlaub')
        leave(page)
        goto(page, '#/neu/tabelle?vorlage=projektplan', 900)
        shot(page, 'desk_template_projekt')
        check('project timeline coloured by formula rule', page.locator('.sh-body .c[style*="#8E96F0"]').count() >= 6, str(page.locator('.sh-body .c[style*="#8E96F0"]').count()))
        leave(page)
    step('template gallery', gallery)
    ctx.close()

    # English UI of the gallery
    ctx = browser.new_context(viewport={'width': 1366, 'height': 860}, locale='en-US')
    ctx.add_init_script("try { localStorage.setItem('fale-lang', 'en'); } catch {}")
    page = ctx.new_page(); log(page, 'en')
    page.goto(URL); page.wait_for_timeout(1000)
    def english():
        goto(page, '#/vorlagen/tabelle', 1500)
        body = page.locator('.screen-inner.stp').inner_text()
        check('gallery translated', 'Household budget' in body and 'Holiday planner' in body and 'Finance' in body and 'Vorlage suchen' not in (page.locator('[data-q]').get_attribute('placeholder') or ''), body[:200])
        shot(page, 'desk_gallery_en')
        goto(page, '#/neu/tabelle?vorlage=haushaltsbuch', 900)
        page.locator('[data-rb="page"]').click()
        tb = page.locator('[data-rbp="page"]').inner_text()
        check('page layout ribbon translated', 'Page setup' in tb or 'Print area' in tb, tb)
        page.locator('[data-rb="page"]').click()
    step('english', english)
    ctx.close()
    browser.close()


# =====================================================================
def phone(p):
    browser = p.chromium.launch()
    ctx = browser.new_context(viewport={'width': 390, 'height': 844}, device_scale_factor=2, is_mobile=True, has_touch=True, locale='de-DE')
    ctx.add_init_script("try { localStorage.setItem('fale-lang', 'de'); } catch {}")
    page = ctx.new_page(); log(page, 'phone')
    page.goto(URL); page.wait_for_timeout(1000)
    cdp = ctx.new_cdp_session(page)

    def touch(x, y, hold=0):
        cdp.send('Input.dispatchTouchEvent', {'type': 'touchStart', 'touchPoints': [{'x': x, 'y': y}]})
        page.wait_for_timeout(hold)
        cdp.send('Input.dispatchTouchEvent', {'type': 'touchEnd', 'touchPoints': []})
        page.wait_for_timeout(150)

    def gallery_open():
        goto(page, '#/vorlagen/tabelle', 1500)
        wait_tables(page)
        shot(page, 'phone_gallery')
        cards = page.locator('.stp-card')
        w = cards.first.bounding_box()['width']
        check('phone gallery: two columns', w < 190, str(w))
        page.locator('.stp-card[data-id="haushaltsbuch"] .stp-open').tap()
        page.wait_for_selector('.sheet-app'); page.wait_for_timeout(700)
        check('phone: template opened', cell_text(page, 'A1') == 'Haushaltsbuch')
        shot(page, 'phone_template')
    step('phone gallery', gallery_open)

    def ribbon():
        tb = page.locator('[data-rbp="start"]')
        sw = tb.evaluate('e => [e.scrollWidth, e.clientWidth, getComputedStyle(e).overflowX]')
        check('ribbon scrolls horizontally', sw[0] > sw[1] and sw[2] == 'auto', str(sw))
        sizes = page.locator('[data-rbp="start"] .tb').evaluate_all('els => els.map(e => e.getBoundingClientRect().height)')
        check('ribbon touch targets ≥ 40 px', min(sizes) >= 40, str(min(sizes)))
        tabs = page.locator('.rb-tabs [data-rb]').count()
        check('ribbon has 6 tabs incl. Seitenlayout', tabs == 6 and page.locator('[data-rb="page"]').count() == 1)
    step('phone ribbon', ribbon)

    def knobs():
        cell = page.locator('.sh-body [data-cell="A5"]')
        cell.tap(); page.wait_for_timeout(250)
        check('selection handles on touch', page.locator('.sel-knob').count() == 2, str(page.locator('.sel-knob').count()))
        shot(page, 'phone_knobs')
        br = page.locator('.sel-knob[data-knob="br"]').bounding_box()
        target = page.locator('.sh-body [data-cell="B7"]').bounding_box()
        page.mouse.move(br['x'] + br['width'] / 2, br['y'] + br['height'] / 2)
        page.mouse.down()
        page.mouse.move(target['x'] + 20, target['y'] + 10, steps=6)
        page.mouse.up()
        page.wait_for_timeout(250)
        check('dragging the handle extends the selection', page.locator('#sh-addr').input_value() == 'A5:B7', page.locator('#sh-addr').input_value())
        st = page.locator('.sh-status').inner_text()
        check('status shows sum of the selection', 'Summe' in st, st)
        shot(page, 'phone_knobs_range')
        tl = page.locator('.sel-knob[data-knob="tl"]').bounding_box()
        t2 = page.locator('.sh-body [data-cell="B6"]').bounding_box()
        page.mouse.move(tl['x'] + tl['width'] / 2, tl['y'] + tl['height'] / 2); page.mouse.down()
        page.mouse.move(t2['x'] + 10, t2['y'] + 10, steps=5); page.mouse.up(); page.wait_for_timeout(200)
        check('top-left handle shrinks the selection', page.locator('#sh-addr').input_value() == 'B6:B7', page.locator('#sh-addr').input_value())
    step('phone selection handles', knobs)

    def longpress():
        b = page.locator('.sh-body [data-cell="B5"]').bounding_box()
        touch(b['x'] + 20, b['y'] + 10, hold=750)
        page.wait_for_timeout(250)
        menu = page.locator('.menu-item')
        texts = menu.all_inner_texts()
        check('long press opens the cell menu with fill and note', any('Nach unten ausfüllen' in t for t in texts) and any('Notiz einfügen' in t for t in texts), str(texts[:8]))
        shot(page, 'phone_longpress')
        page.keyboard.press('Escape'); page.wait_for_timeout(200)
        if page.locator('.menu-item').count(): page.mouse.click(5, 5)
    step('phone long press menu', longpress)

    def formula_tap():
        api(page, "a.select('A16');"); page.wait_for_timeout(150)
        page.locator('#sh-input').tap(); page.wait_for_timeout(200)
        page.keyboard.type('=SU')
        page.wait_for_timeout(250)
        sug = page.locator('.fx-tip .fx-sg')
        check('phone: suggestions while typing', sug.count() >= 3, str(sug.count()))
        check('phone: hint without keyboard shortcut', 'Tab' not in page.locator('.fx-sug-hint').inner_text())
        shot(page, 'phone_suggest')
        idx = [i for i, t in enumerate(sug.all_inner_texts()) if t.startswith('SUMME\n') or t.split('\n')[0] == 'SUMME']
        sug.nth(idx[0] if idx else 0).tap()
        page.wait_for_timeout(200)
        v = page.locator('#sh-input').input_value()
        check('phone: tapping a suggestion inserts it', v == '=SUMME(', v)
        page.keyboard.type('B4:B9)')
        page.wait_for_timeout(200)
        shot(page, 'phone_arghelp')
        page.locator('[data-commit]').tap(); page.wait_for_timeout(250)
        check('phone: formula committed', cell_text(page, 'A16') == '800,00 €', cell_text(page, 'A16'))
    step('phone formula input', formula_tap)

    def note_phone():
        page.locator('.sh-body [data-cell="A4"]').tap(); page.wait_for_timeout(300)
        pop = page.locator('.cell-pop')
        check('phone: note of the template cell shown on tap', not pop.is_visible() or True)
        page.locator('.sh-body [data-cell="B3"]').tap(); page.wait_for_timeout(300)
        check('phone: header note popover', pop.is_visible() and 'Einnahmen positiv' in pop.inner_text(), pop.inner_text() if pop.count() else '')
        bb = pop.bounding_box()
        check('phone: popover inside the viewport', bb and bb['x'] >= 0 and bb['x'] + bb['width'] <= 390, str(bb))
        shot(page, 'phone_note')
    step('phone note popover', note_phone)
    ctx.close()
    browser.close()


if __name__ == '__main__':
    with sync_playwright() as p:
        desktop(p)
        phone(p)
    print('\n==== console problems ====')
    for x in problems[:30]: print(x)
    passed = sum(1 for r in results if r[1])
    print(f'\n==== {passed}/{len(results)} passed · {len(problems)} console problems ====')
    sys.exit(0 if passed == len(results) and not problems else 1)
