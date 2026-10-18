"""Office 3: XLSX fidelity. Every spreadsheet template and a feature workbook are written by our writer (in the browser),
checked with openpyxl (structure), recalculated by LibreOffice headless (values compared with our engine), converted to PDF
by LibreOffice, and the LibreOffice re-save is read back by our reader. Server: OFFICE_URL (default http://127.0.0.1:8850/)."""
import os, sys, json, base64, subprocess, shutil, math, traceback
from playwright.sync_api import sync_playwright
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
WORK = os.path.join(HERE, 'work_rt'); shutil.rmtree(WORK, ignore_errors=True); os.makedirs(WORK)
LO_OUT = os.path.join(WORK, 'lo'); os.makedirs(LO_OUT)
PROFILE = os.path.join(HERE, 'lo_profile')
URL = os.environ.get('OFFICE_URL', 'http://127.0.0.1:8850/')
results = []
def check(name, cond, info=''):
    results.append((name, bool(cond), info))
    print(('PASS ' if cond else 'FAIL ') + name + (f'  [{info}]' if info and not cond else ''), flush=True)

GEN = r"""async () => {
  const T = await import('/js/office/sheet-templates.js'); const W = await import('/js/office/workbook.js');
  const FX = await import('/js/office/sheetfx.js'); const X = await import('/js/office/xlsx-write.js'); const F = await import('/js/office/formula.js');
  const b64 = async blob => { const buf = new Uint8Array(await blob.arrayBuffer()); let s = ''; for (let i = 0; i < buf.length; i += 0x8000) s += String.fromCharCode(...buf.subarray(i, i + 0x8000)); return btoa(s); };
  const dump = wb => wb.sheets.map(sh => { const vals = {}; for (const [k, c] of sh.cells) if (c.f) { const v = c.cv; vals[F.addr(W.R_OF(k), W.C_OF(k))] = v && v.code ? { err: v.code } : v; }
    return { name: sh.name, vals, cf: sh.cf.length, dv: sh.dv.length, notes: sh.notes.size, links: sh.links.size, freeze: sh.freeze, freezeC: sh.freezeC, page: sh.page || null, charts: sh.charts.length, outline: [...sh.outline], merges: sh.merges.length }; });
  const out = [];
  for (const t of T.SHEET_TEMPLATE_DEFS) {
    const wb = T.buildSheetTemplate(t.id);
    wb.sheets.forEach(s => { FX.ensureSheetExtras(s); W.ensureSheetModel(s); }); W.recalc(wb);
    out.push({ id: t.id, data: await b64(await X.writeXlsx(wb, { title: t.title })), sheets: dump(wb) });
  }
  // feature workbook
  const wb = W.newWorkbook(); const sh = wb.sheets[0]; sh.name = 'Funktionen';
  const put = (a, text) => { const p = F.parseAddr(a); const x = W.parseInput(text); const cell = {}; if (x.fmt) cell.s = { fmt: x.fmt }; if (x.f) { cell.f = x.f; cell.da = true; } else cell.v = x.v; W.setCell(sh, p.r, p.c, cell); };
  [['A1', 'Wert'], ['A2', '10'], ['A3', '-4'], ['A4', '25'], ['A5', '7'], ['A6', '25'], ['B1', 'Text'], ['B2', 'Apfel'], ['B3', 'Birne'], ['B4', 'Apfel'], ['B5', ''], ['B6', 'Kiwi'],
   ['C2', '=SUMME(A2:A6)'], ['C3', '=SUMMEWENN(B2:B6;"Apfel";A2:A6)'], ['C4', '=MITTELWERT(A2:A6)'], ['C5', '=STEIGUNG(A2:A6;{1;2;3;4;5})'], ['C6', '=KUMZINSZ(5%/12;60;20000;1;12;0)'],
   ['D2', '=GDA(2400;300;10;1)'], ['D3', '=DIA(30000;7500;10;1)'], ['D4', '=QIKV({-120000;39000;30000;21000;37000;46000};0,1;0,12)'], ['D5', '=NORM.VERT(42;40;1,5;WAHR)'], ['D6', '=BINOM.VERT(6;10;0,5;FALSCH)'],
   ['E2', '=RÖMISCH(2026)'], ['E3', '=ARABISCH("MMXXVI")'], ['E4', '=DEZINBIN(9;8)'], ['E5', '=HEXINDEZ("FF")'], ['E6', '=MDET({3.6.1;1.1.0;3.10.2})'],
   ['F2', '=KURT(3;4;5;2;3;4;5;6;4;7)'], ['F3', '=SCHIEFE(3;4;5;2;3;4;5;6;4;7)'], ['F4', '=GESTUTZTMITTEL({4.5.6.7.2.3.4.5.1.2.3};0,2)'], ['F5', '=QUANTIL.EXKL({1.2.3.6.6.6.7.8.9};0,25)'], ['F6', '=KOVARIANZ.S({3.2.4.5.6};{9.7.12.15.17})'],
   ['G2', '=DBSUMME(A1:B6;"Wert";H1:H2)'], ['H1', 'Text'], ['H2', 'Apfel'], ['G3', '=AGGREGAT(9;6;A2:A6)'], ['G4', '=SUMMEX2MY2({2.3.9};{6.5.11})'], ['G5', '=ZW2(1;{0,09.0,11.0,1})'], ['G6', '=PDURATION(2,5%;2000;2200)'],
   ['I2', '=LAMBDA(x;x*2)(21)'], ['I3', '=TEXTVERKETTEN("-";WAHR;B2:B4)'], ['I4', '=REGEXTEST("abc123";"[0-9]+")'],
  ].forEach(([a, t]) => put(a, t));
  sh.notes.set(W.KEY(0, 0), { text: 'Notiz mit Umlauten: äöü ß' }); sh.notes.set(W.KEY(3, 1), { text: 'Zweite\nZeile' });
  sh.links.set(W.KEY(1, 1), { url: 'https://example.org/apfel?x=1&y=2' }); sh.links.set(W.KEY(2, 1), { ref: "'Zweites Blatt'!A1" }); sh.links.set(W.KEY(5, 1), { url: 'mailto:info@example.org' });
  const rules = [{ type: 'gt', v1: '5', style: 'green' }, { type: 'ge', v1: '10', style: 'yellow' }, { type: 'le', v1: '0', style: 'red' }, { type: 'ne', v1: '7', style: 'blue' }, { type: 'top', v1: '2', style: 'red' }, { type: 'bottom', v1: '20', percent: true, style: 'gray' },
    { type: 'aboveAvg', style: 'green' }, { type: 'belowAvg', style: 'red' }, { type: 'dup', style: 'yellow' }, { type: 'unique', style: 'blue' }, { type: 'blank', style: 'gray' }, { type: 'notblank', style: 'bold' }, { type: 'error', style: 'red' },
    { type: 'formula', f: 'AND($A2>0,MOD(ROW(),2)=0)', style: 'done' }, { type: 'bar', bar: 'green' }, { type: 'icons', icons: 'arrows' }, { type: 'scale', scale: 'gr' }, { type: 'contains', v1: 'Ap', style: 'red' }, { type: 'between', v1: '1', v2: '9', style: 'blue' }];
  sh.cf = rules.map((r, i) => ({ id: 'r' + i, range: 'A2:A6', ...r }));
  sh.dv = [{ range: 'B2:B6', type: 'list', list: ['Apfel', 'Birne', 'Kiwi'] }, { range: 'A2:A6', type: 'whole', op: 'between', v1: '-100', v2: '100', list: [] }];
  sh.page = { orient: 'landscape', fit: 'width', grid: true, head: true, area: 'A1:I6', titleRows: { r1: 0, r2: 0 }, margins: 'narrow', center: true, footer: true };
  sh.outline = new Map([[2, 1], [3, 1], [4, 2]]); sh.hiddenRows = new Set([4]);
  sh.merges.push({ r1: 7, c1: 0, r2: 7, c2: 3 }); W.setCell(sh, 7, 0, { v: 'Verbunden', s: { b: true, align: 'center' } });
  sh.freeze = 1; sh.freezeC = 1; sh.colW[1] = 140;
  const sh2 = W.newSheet('Zweites Blatt'); wb.sheets.push(sh2);
  [['Ziel', 'Umsatz', 'Kosten'], ['Q1', 120, 80], ['Q2', 150, 95], ['Q3', 170, 90]].forEach((row, r) => row.forEach((v, c) => W.setCell(sh2, r, c, { v })));
  sh2.charts = [{ id: 'c1', type: 'combo', range: 'A1:C4', title: 'Kombi', xTitle: '', yTitle: '', legend: 'b', palette: 'standard', byRows: null, series: null, colors: null, x: 300, y: 0, w: 400, h: 260 },
    { id: 'c2', type: 'doughnut', range: 'A1:B4', title: 'Ring', xTitle: '', yTitle: '', legend: 'r', palette: 'standard', byRows: null, series: null, colors: null, x: 300, y: 300, w: 400, h: 260 }];
  wb.sheets.forEach(s => { FX.ensureSheetExtras(s); W.ensureSheetModel(s); }); W.recalc(wb);
  out.push({ id: 'features', data: await b64(await X.writeXlsx(wb, { title: 'Features' })), sheets: dump(wb) });
  return out;
}"""

READ = r"""async (b64s) => {
  const R = await import('/js/office/xlsx-read.js'); const F = await import('/js/office/formula.js'); const W = await import('/js/office/workbook.js');
  const out = {};
  for (const [name, b64] of Object.entries(b64s)) {
    const bin = atob(b64); const buf = new Uint8Array(bin.length); for (let i = 0; i < bin.length; i++) buf[i] = bin.charCodeAt(i);
    try {
      const wb = await R.readXlsx(buf.buffer, name);
      out[name] = wb.sheets.map(sh => ({ name: sh.name, chartTypes: sh.charts.map(c => c.type), cf: sh.cf.map(r => r.type), dv: sh.dv.length, notes: [...sh.notes.values()].map(n => n.text), links: [...sh.links.values()],
        page: sh.page || null, outline: [...sh.outline], freeze: sh.freeze, freezeC: sh.freezeC, merges: sh.merges.length, charts: sh.charts.length,
        vals: Object.fromEntries([...sh.cells].filter(([k, c]) => c.f).slice(0, 4000).map(([k, c]) => [F.addr(W.R_OF(k), W.C_OF(k)), c.cv && c.cv.code ? { err: c.cv.code } : c.cv])) }));
    } catch (e) { out[name] = { error: String(e && e.message || e) }; }
  }
  return out;
}"""

import datetime
EPOCH = datetime.datetime(1899, 12, 30)
def norm(v):
    if isinstance(v, datetime.datetime): return (v - EPOCH).total_seconds() / 86400
    if isinstance(v, datetime.date): return (datetime.datetime(v.year, v.month, v.day) - EPOCH).days
    if isinstance(v, datetime.time): return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    if isinstance(v, datetime.timedelta): return v.total_seconds() / 86400
    return v
def same(a, b):
    a, b = norm(a), norm(b)
    if isinstance(a, bool) or isinstance(b, bool): return a == b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= max(1e-6, abs(b) * 1e-7)
    if a in (None, '') and b in (None, ''): return True
    return str(a) == str(b)

with sync_playwright() as p:
    browser = p.chromium.launch()
    page = browser.new_page()
    page.on('pageerror', lambda e: print('PAGEERR', e))
    page.goto(URL); page.wait_for_timeout(800)
    gen = page.evaluate(GEN)
    files = {}
    for g in gen:
        path = os.path.join(WORK, f"{g['id']}.xlsx")
        open(path, 'wb').write(base64.b64decode(g['data']))
        files[g['id']] = (path, g['sheets'])
    n_defs = page.evaluate("async () => (await import('/js/office/sheet-templates.js')).SHEET_TEMPLATE_DEFS.length")
    check('all templates written', {g['id'] for g in gen} >= set(page.evaluate("async () => (await import('/js/office/sheet-templates.js')).SHEET_TEMPLATE_DEFS.map(t => t.id)")) and n_defs >= 40, f'{len(gen)} of {n_defs}')

    # ---- structure via openpyxl
    for tid, (path, sheets) in files.items():
        try:
            wb = openpyxl.load_workbook(path)
        except Exception as e:
            check(f'{tid}: openpyxl opens', False, str(e)); continue
        check(f'{tid}: sheet names', wb.sheetnames == [s['name'] for s in sheets], str(wb.sheetnames))
        for s in sheets:
            ws = wb[s['name']]
            nf = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('='))
            if s['vals']: check(f"{tid}/{s['name']}: formulas kept", nf >= len(s['vals']) * 0.99, f"{nf} vs {len(s['vals'])}")
            ncf = sum(len(r.rules) for r in ws.conditional_formatting)
            check(f"{tid}/{s['name']}: conditional formats", ncf == s['cf'], f"{ncf} vs {s['cf']}")
            check(f"{tid}/{s['name']}: validations", len(ws.data_validations.dataValidation) == s['dv'], f"{len(ws.data_validations.dataValidation)} vs {s['dv']}")
            ncom = sum(1 for row in ws.iter_rows() for c in row if c.comment)
            check(f"{tid}/{s['name']}: notes as comments", ncom == s['notes'], f"{ncom} vs {s['notes']}")
            nl = sum(1 for row in ws.iter_rows() for c in row if c.hyperlink)
            check(f"{tid}/{s['name']}: hyperlinks", nl == s['links'], f"{nl} vs {s['links']}")
            if s['freeze'] or s['freezeC']:
                exp = f"{openpyxl.utils.get_column_letter(s['freezeC'] + 1)}{s['freeze'] + 1}"
                check(f"{tid}/{s['name']}: freeze panes", ws.freeze_panes == exp, f"{ws.freeze_panes} vs {exp}")
            pg = s['page']
            if pg and pg.get('titleRows'):
                tr = pg['titleRows']
                check(f"{tid}/{s['name']}: print titles", str(ws.print_title_rows).replace('$', '') == f"{tr['r1'] + 1}:{tr['r2'] + 1}", str(ws.print_title_rows))
            if pg and pg.get('orient') in ('portrait', 'landscape'):
                check(f"{tid}/{s['name']}: orientation", ws.page_setup.orientation == pg['orient'], str(ws.page_setup.orientation))
            if pg and pg.get('fit') == 'width':
                check(f"{tid}/{s['name']}: fit to width", ws.sheet_properties.pageSetUpPr is not None and ws.sheet_properties.pageSetUpPr.fitToPage and str(ws.page_setup.fitToWidth) == '1' and str(ws.page_setup.fitToHeight) == '0', f"{ws.page_setup.fitToWidth}/{ws.page_setup.fitToHeight}")
            if pg and pg.get('area'):
                check(f"{tid}/{s['name']}: print area", ws.print_area is not None and pg['area'].replace(':', '') in ws.print_area.replace('$', '').replace(':', ''), str(ws.print_area))
            if s['outline']:
                lv = {r: ws.row_dimensions[r + 1].outline_level for r, _ in s['outline']}
                check(f"{tid}/{s['name']}: outline levels", all(lv[r] == l for r, l in s['outline']), str(lv))

    # ---- LibreOffice: recalculate on load, re-save as xlsx, compare values
    lo = ['soffice', f'-env:UserInstallation=file://{PROFILE}', '--headless', '--calc']
    ids = list(files.keys())
    r = subprocess.run(lo + ['--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', LO_OUT] + [files[i][0] for i in ids], capture_output=True, timeout=600)
    check('LibreOffice converts all files', all(os.path.exists(os.path.join(LO_OUT, f'{i}.xlsx')) for i in ids), r.stderr.decode()[-300:])
    skip_fn = ('LAMBDA', 'REGEXTEST', 'AGGREGATE', 'PDURATION', 'DPRODUCT')
    total = 0; mism = []
    for tid in ids:
        out = os.path.join(LO_OUT, f'{tid}.xlsx')
        if not os.path.exists(out): continue
        wbv = openpyxl.load_workbook(out, data_only=True)
        wbf = openpyxl.load_workbook(files[tid][0])
        for s in files[tid][1]:
            if s['name'] not in wbv.sheetnames: mism.append(f'{tid}: sheet {s["name"]} missing'); continue
            ws, wf = wbv[s['name']], wbf[s['name']]
            for a, v in s['vals'].items():
                f = wf[a].value
                if isinstance(f, str) and any(x in f.upper() for x in skip_fn): continue
                lv = ws[a].value
                if isinstance(v, dict):  # error values
                    total += 1
                    if not (isinstance(lv, str) and lv.startswith('#')): mism.append(f'{tid}!{a}: ours {v} LO {lv!r} [{f}]')
                    continue
                if v is None: continue
                total += 1
                if isinstance(v, str) and v == '' and lv in (None, ''): continue
                if not same(lv, v): mism.append(f'{tid}/{s["name"]}!{a}: ours {v!r} LO {lv!r} [{f}]')
    check(f'LibreOffice recalculation matches our values ({total} formula cells)', not mism, ' | '.join(mism[:12]) + (f' … +{len(mism) - 12}' if len(mism) > 12 else ''))
    print(f'compared {total} formula results with LibreOffice, {len(mism)} differences')

    # ---- LibreOffice: PDF export of every file (opens without errors, has pages)
    pdfdir = os.path.join(WORK, 'pdf'); os.makedirs(pdfdir)
    r = subprocess.run(lo + ['--convert-to', 'pdf', '--outdir', pdfdir] + [files[i][0] for i in ids], capture_output=True, timeout=600)
    ok_pdf = [i for i in ids if os.path.exists(os.path.join(pdfdir, f'{i}.pdf')) and os.path.getsize(os.path.join(pdfdir, f'{i}.pdf')) > 1000]
    check('LibreOffice exports every file to PDF', len(ok_pdf) == len(ids), str(set(ids) - set(ok_pdf)))

    # ---- our reader: our own files and the LibreOffice re-saves
    b64s = {f'own_{i}': base64.b64encode(open(files[i][0], 'rb').read()).decode() for i in ids}
    b64s.update({f'lo_{i}': base64.b64encode(open(os.path.join(LO_OUT, f'{i}.xlsx'), 'rb').read()).decode() for i in ids if os.path.exists(os.path.join(LO_OUT, f'{i}.xlsx'))})
    back = page.evaluate(READ, b64s)
    for i in ids:
        own = back.get(f'own_{i}')
        if not isinstance(own, list): check(f'{i}: our reader opens own file', False, str(own)); continue
        for s, o in zip(files[i][1], own):
            check(f'{i}/{s["name"]}: own round trip (cf/dv/notes/links/merges/charts)', len(o['cf']) == s['cf'] and o['dv'] == s['dv'] and len(o['notes']) == s['notes'] and len(o['links']) == s['links'] and o['merges'] == s['merges'] and o['charts'] == s['charts'],
                  f"cf {len(o['cf'])}/{s['cf']} dv {o['dv']}/{s['dv']} notes {len(o['notes'])}/{s['notes']} links {len(o['links'])}/{s['links']} charts {o['charts']}/{s['charts']}")
            diff = [a for a, v in s['vals'].items() if not (isinstance(v, dict) or same(o['vals'].get(a), v))]
            check(f'{i}/{s["name"]}: own round trip values', not diff, ' '.join(f"{a}:{s['vals'][a]!r}->{o['vals'].get(a)!r}" for a in diff[:6]))
            if s['page']: check(f'{i}/{s["name"]}: page setup read back', bool(o['page']) and o['page'].get('orient') == s['page'].get('orient') and (o['page'].get('titleRows') == s['page'].get('titleRows')), f"{o['page']} vs {s['page']}")
        lo_back = back.get(f'lo_{i}')
        if isinstance(lo_back, list):
            s0, o0 = files[i][1][0], lo_back[0]
            check(f'{i}: LibreOffice re-save readable (notes/links/validations kept)', len(o0['notes']) == s0['notes'] and len(o0['links']) == s0['links'] and o0['dv'] >= min(1, s0['dv']), f"notes {len(o0['notes'])}/{s0['notes']} links {len(o0['links'])}/{s0['links']} dv {o0['dv']}/{s0['dv']}")
        else: check(f'{i}: LibreOffice re-save readable', False, str(lo_back))
    # feature workbook details
    f0 = back['own_features'][0]
    check('features: all rule types read back', sorted(f0['cf']) == sorted(['gt', 'ge', 'le', 'ne', 'top', 'bottom', 'aboveAvg', 'belowAvg', 'dup', 'unique', 'blank', 'notblank', 'error', 'formula', 'bar', 'icons', 'scale', 'contains', 'between']), str(f0['cf']))
    check('features: notes text kept', f0['notes'] == ['Notiz mit Umlauten: äöü ß', 'Zweite\nZeile'], str(f0['notes']))
    check('features: links kept (web, internal, mail)', {'url': 'https://example.org/apfel?x=1&y=2'} in f0['links'] and {'ref': "'Zweites Blatt'!A1"} in f0['links'] and {'url': 'mailto:info@example.org'} in f0['links'], str(f0['links']))
    check('features: page setup kept', f0['page'] and f0['page']['orient'] == 'landscape' and f0['page']['fit'] == 'width' and f0['page']['area'] == 'A1:I6' and f0['page']['titleRows'] == {'r1': 0, 'r2': 0} and f0['page']['margins'] == 'narrow' and f0['page']['center'] and f0['page']['footer'], str(f0['page']))
    check('features: outline kept', sorted(map(tuple, f0['outline'])) == [(2, 1), (3, 1), (4, 2)], str(f0['outline']))
    check('features: new functions survive the round trip', f0['vals'].get('I2') == 42 and f0['vals'].get('E2') == 'MMXXVI' and abs(f0['vals'].get('D5', 0) - 0.90878878) < 1e-6, str({k: f0['vals'].get(k) for k in ('I2', 'E2', 'D5')}))
    f1 = back['own_features'][1]
    check('features: combo and doughnut charts round trip', f1['chartTypes'] == ['combo', 'doughnut'], str(f1['chartTypes']))
    lof = back.get('lo_features')
    if isinstance(lof, list) and len(lof) > 1: check('features: LibreOffice keeps combo/doughnut charts', sorted(lof[1]['chartTypes']) == ['combo', 'doughnut'], str(lof[1]['chartTypes']))
    if isinstance(lof, list):
        check('features: LibreOffice keeps conditional formats (types)', len(lof[0]['cf']) >= 15, str(lof[0]['cf']))
        check('features: LibreOffice keeps notes/links/outline', len(lof[0]['notes']) == 2 and len(lof[0]['links']) >= 2 and len(lof[0]['outline']) == 3, f"{lof[0]['notes']} {lof[0]['links']} {lof[0]['outline']}")
    browser.close()

passed = sum(1 for r in results if r[1])
print(f'\n==== {passed}/{len(results)} passed ====')
sys.exit(0 if passed == len(results) else 1)
