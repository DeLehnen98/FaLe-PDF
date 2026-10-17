"""Create sample office files for test_office.py."""
import os, io, datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PRGB
from PIL import Image

D = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'officedata')
os.makedirs(D, exist_ok=True)

# small image
img = Image.new('RGB', (240, 120), (59, 55, 160))
for x in range(240):
    for y in range(40, 80):
        img.putpixel((x, y), (45, 212, 191))
img.save(os.path.join(D, 'bild.png'))

# ---- docx
doc = Document()
doc.add_heading('Projektbericht Zitronenfalter', level=1)
p = doc.add_paragraph('Dies ist ein ')
r = p.add_run('fetter'); r.bold = True
p.add_run(' und ')
r = p.add_run('kursiver'); r.italic = True
p.add_run(' Text mit ')
r = p.add_run('roter Farbe'); r.font.color.rgb = RGBColor(0xC0, 0x10, 0x10)
p.add_run(' sowie Umlauten: Größe, Übermaß, Ärger – 12 € und ß.')
c = doc.add_paragraph('Zentrierter Absatz'); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading('Aufzählung', level=2)
for t in ['Erster Punkt', 'Zweiter Punkt', 'Dritter Punkt']:
    doc.add_paragraph(t, style='List Bullet')
doc.add_heading('Nummerierung', level=2)
for t in ['Schritt eins', 'Schritt zwei']:
    doc.add_paragraph(t, style='List Number')
doc.add_heading('Tabelle', level=2)
t = doc.add_table(rows=3, cols=3); t.style = 'Table Grid'
for i, row in enumerate([['Artikel', 'Menge', 'Preis'], ['Äpfel', '3', '1,99 €'], ['Birnen', '5', '2,49 €']]):
    for j, v in enumerate(row):
        t.cell(i, j).text = v
doc.add_paragraph('Bild:')
doc.add_picture(os.path.join(D, 'bild.png'), width=Inches(2))
doc.add_page_break()
doc.add_paragraph('Zweite Seite mit dem Wort Seepferdchen.')
long = ' '.join(['Langer Fließtext zum Testen des Zeilenumbruchs.'] * 40)
doc.add_paragraph(long)
doc.save(os.path.join(D, 'bericht.docx'))

# ---- xlsx
wb = Workbook()
ws = wb.active
ws.title = 'Umsatz'
ws.append(['Monat', 'Einnahmen', 'Ausgaben', 'Saldo'])
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill('solid', fgColor='E9E8FB')
data = [('Januar', 1200, 800), ('Februar', 1500.5, 950.25), ('März', 980, 1100)]
for i, (m, e, a) in enumerate(data, start=2):
    ws.append([m, e, a, f'=B{i}-C{i}'])
ws['A6'] = 'Summe'
ws['B6'] = '=SUM(B2:B4)'
ws['C6'] = '=SUM(C2:C4)'
ws['D6'] = '=SUM(D2:D4)'
ws['A7'] = 'Durchschnitt'
ws['B7'] = '=AVERAGE(B2:B4)'
ws['A8'] = 'Positiv'
ws['B8'] = '=COUNTIF(D2:D4,">0")'
ws['A9'] = 'Status'
ws['B9'] = '=IF(D6>0,"Gewinn","Verlust")'
ws['A10'] = 'Datum'
ws['B10'] = datetime.date(2026, 3, 15)
ws['B10'].number_format = 'DD.MM.YYYY'
ws['A11'] = 'Quote'
ws['B11'] = 0.256
ws['B11'].number_format = '0.0%'
ws['A12'] = 'Wahr'
ws['B12'] = True
ws.column_dimensions['A'].width = 16
ws.merge_cells('F1:G1')
ws['F1'] = 'Zusammengeführt'
ws2 = wb.create_sheet('Preise')
ws2.append(['Artikel', 'Preis'])
ws2.append(['Äpfel', 1.99])
ws2.append(['Birnen', 2.49])
ws2.append(['Kirschen', 4.2])
ws2['D1'] = 'Suche'
ws2['D2'] = '=VLOOKUP("Birnen",A2:B4,2,FALSE)'
ws2['D3'] = '=Umsatz!B6*2'
ws2['D4'] = '=ROUND(B4*1.19,2)'
ws2['D5'] = '=CONCATENATE(A2," & ",A3)'
ws2['D6'] = '=SUMIF(B2:B4,">2")'
ws2['D7'] = '=COUNTA(A1:A10)'
ws2['D8'] = '=MAX(B2:B4)-MIN(B2:B4)'
wb.save(os.path.join(D, 'umsatz.xlsx'))

# ---- csv (German)
with open(os.path.join(D, 'liste.csv'), 'w', encoding='utf-8-sig', newline='') as f:
    f.write('Name;Betrag;Datum\r\nMüller;1.234,50;01.02.2026\r\nSchäfer;99,90;15.03.2026\r\n"Weiß; Anna";12;31.12.2025\r\n')

# ---- pptx
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = 'Quartalsbericht Seepferdchen'
s.placeholders[1].text = 'Untertitel mit Umlauten äöü'
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = 'Agenda'
tf = s.placeholders[1].text_frame
tf.text = 'Erster Punkt'
for t in ['Zweiter Punkt', 'Dritter Punkt']:
    p = tf.add_paragraph(); p.text = t
p = tf.add_paragraph(); p.text = 'Unterpunkt'; p.level = 1
s = prs.slides.add_slide(prs.slide_layouts[6])
box = s.shapes.add_shape(1, PI(1), PI(1), PI(3), PI(1.5))
box.fill.solid(); box.fill.fore_color.rgb = PRGB(0x0B, 0x8F, 0x80)
box.text_frame.text = 'Grüne Box'
tb = s.shapes.add_textbox(PI(5), PI(1), PI(4), PI(1))
tb.text_frame.text = 'Freies Textfeld'
tb.text_frame.paragraphs[0].runs[0].font.size = PP(28)
tb.text_frame.paragraphs[0].runs[0].font.bold = True
s.shapes.add_picture(os.path.join(D, 'bild.png'), PI(1), PI(3.5), PI(3))
s.background.fill.solid(); s.background.fill.fore_color.rgb = PRGB(0xFF, 0xF6, 0xF0)
prs.save(os.path.join(D, 'praesentation.pptx'))

with open(os.path.join(D, 'notiz.txt'), 'w', encoding='utf-8') as f:
    f.write('Einkaufsliste\nÄpfel\nBrötchen für 3,50 €\n\nEnde der Notiz')
with open(os.path.join(D, 'readme.md'), 'w', encoding='utf-8') as f:
    f.write('# Überschrift\n\nEin **fetter** Satz mit *kursiv*.\n\n- Punkt A\n- Punkt B\n\n1. Eins\n2. Zwei\n')
print('ok', os.listdir(D))

# ---- PDF with a table (for PDF → Excel)
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas as rlcanvas
pdfp = os.path.join(D, 'tabelle.pdf')
cv = rlcanvas.Canvas(pdfp, pagesize=A4)
cv.setFont('Helvetica-Bold', 16)
cv.drawString(60, 790, 'Preisliste Sommer')
cv.setFont('Helvetica', 10)
cv.drawString(60, 770, 'Alle Preise inklusive Mehrwertsteuer.')
rows = [('Artikel', 'Menge', 'Einzelpreis', 'Gesamt'), ('Sonnenschirm', '2', '49,90 €', '99,80 €'), ('Liegestuhl', '4', '1.234,50 €', '4.938,00 €'), ('Kühlbox', '1', '35,00 €', '35,00 €'), ('Strandtuch', '10', '9,99 €', '99,90 €')]
y = 730
for i, r in enumerate(rows):
    cv.setFont('Helvetica-Bold' if i == 0 else 'Helvetica', 11)
    cv.drawString(60, y, r[0])
    cv.drawRightString(260, y, r[1])
    cv.drawRightString(390, y, r[2])
    cv.drawRightString(520, y, r[3])
    y -= 22
cv.showPage()
cv.setFont('Helvetica', 11)
cv.drawString(60, 790, 'Seite zwei')
y = 760
for i, r in enumerate([('Monat', 'Umsatz'), ('Juni', '12.500'), ('Juli', '14.250')]):
    cv.drawString(60, y, r[0]); cv.drawRightString(300, y, r[1]); y -= 22
cv.save()
print('pdf ok')
